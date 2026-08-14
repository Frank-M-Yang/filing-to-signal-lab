"""Create the locked preliminary effectiveness research memo and chart pack.

This reporting script intentionally reads immutable, versioned research outputs. It
does not recompute signals, returns, confidence intervals, or verdicts.
"""

from __future__ import annotations

import json
from collections.abc import Iterable
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
DIAGNOSTIC_DIR = (
    ROOT
    / "data"
    / "processed"
    / "preliminary_effectiveness_diagnostics"
    / "preliminary-effectiveness-diagnostic-v1-ec29d989f4f879f1"
)
PREREG_DIR = (
    ROOT
    / "data"
    / "processed"
    / "preliminary_effectiveness"
    / "preliminary-effectiveness-preregistration-v1-97df4b07e54937e8"
)
LEDGER_DIR = (
    ROOT
    / "data"
    / "processed"
    / "longitudinal_event_returns"
    / "longitudinal-event-return-ledger-v1-295f4ff0fa9a202e"
)
REPAIR_DIR = (
    ROOT
    / "data"
    / "processed"
    / "ticker_transition_repair"
    / "ticker-transition-repair-v1-34e561bf1d728e67"
)
OUTPUT_DIR = ROOT / "outputs" / "reports" / "preliminary_effectiveness_v1"
WORK_DIR = ROOT / "tmp" / "preliminary_effectiveness_report_v1"
FIGURE_DIR = WORK_DIR / "figures"
DOCX_PATH = OUTPUT_DIR / "财报信号初步有效性报告_v1.docx"

FONT_REGULAR = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")

NAVY = "0B2545"
BLUE = "2E74B5"
TEAL = "16817A"
GOLD = "B06B00"
RED = "9B1C1C"
INK = "172B2B"
SLATE = "5E6F78"
MIST = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F4F2"
PALE_GOLD = "FFF4D6"
PALE_RED = "FBE9E7"
WHITE = "FFFFFF"
GRID = "D8E0E5"

PIL_COLORS = {
    "navy": "#0B2545",
    "blue": "#2E74B5",
    "teal": "#16817A",
    "gold": "#B06B00",
    "red": "#9B1C1C",
    "ink": "#172B2B",
    "slate": "#5E6F78",
    "mist": "#F2F4F7",
    "grid": "#D8E0E5",
    "white": "#FFFFFF",
}


def _read_json(path: Path) -> dict[str, Any]:
    """Read a UTF-8 JSON object from an immutable research artifact."""
    return json.loads(path.read_text(encoding="utf-8"))


def _font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    """Return the report's consistent CJK-capable chart font."""
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size=size)


def _pct(value: float, digits: int = 2) -> str:
    """Format a decimal return as a percentage."""
    return f"{value * 100:.{digits}f}%"


def _set_repeat_table_header(row: Any) -> None:
    """Mark a Word table row as a repeating header."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_cell_shading(cell: Any, fill: str) -> None:
    """Apply a solid fill to a Word table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(
    cell: Any,
    *,
    top: int = 80,
    start: int = 120,
    bottom: int = 80,
    end: int = 120,
) -> None:
    """Set Word table cell margins in twips."""
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_border(cell: Any, **edges: dict[str, str]) -> None:
    """Set selected Word table-cell borders."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, edge_data in edges.items():
        edge = tc_borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            tc_borders.append(edge)
        for key in ["val", "sz", "space", "color"]:
            if key in edge_data:
                edge.set(qn(f"w:{key}"), str(edge_data[key]))


def _set_run_font(run: Any, name: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    """Set Latin and East Asian font names for a Word run."""
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _set_style_font(style: Any, name: str, east_asia: str) -> None:
    """Set Latin and East Asian typefaces on a Word style."""
    style.font.name = name
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _keep_paragraph_with_next(paragraph: Any) -> None:
    """Keep a paragraph with the following paragraph or table."""
    paragraph.paragraph_format.keep_with_next = True


def _set_keep_table_rows(table: Any) -> None:
    """Keep each table row from splitting across pages."""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def _add_page_number(paragraph: Any) -> None:
    """Insert a live Word PAGE field."""
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def _populate_running_header(header: Any) -> None:
    """Populate one odd/even Word header with the report identifier."""
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("FILING-TO-SIGNAL LAB  |  PRELIMINARY EFFECTIVENESS")
    _set_run_font(run)
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(SLATE)


def _populate_running_footer(footer: Any) -> None:
    """Populate one odd/even Word footer with disclaimer and page number."""
    footer.is_linked_to_previous = False
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.7))
    footer_table.columns[0].width = Inches(5.8)
    footer_table.columns[1].width = Inches(0.9)
    left = footer_table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left_run = left.add_run("教育研究演示｜不构成投资建议")
    _set_run_font(left_run)
    left_run.font.size = Pt(7.5)
    left_run.font.color.rgb = RGBColor.from_string(SLATE)
    right = footer_table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    page_run = right.add_run("P. ")
    _set_run_font(page_run)
    page_run.font.size = Pt(7.5)
    page_run.font.color.rgb = RGBColor.from_string(SLATE)
    _add_page_number(right)
    for cell in footer_table.row_cells(0):
        _set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
        _set_cell_border(
            cell,
            top={"val": "single", "sz": "6", "space": "0", "color": GRID},
        )


def _configure_document(document: Document) -> None:
    """Apply the standard_business_brief design preset and CJK fallback."""
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri", "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    style_specs = {
        "Title": (26, NAVY, 0, 8),
        "Subtitle": (11.5, SLATE, 0, 8),
        "Heading 1": (16, BLUE, 14, 7),
        "Heading 2": (13, BLUE, 10, 5),
        "Heading 3": (11.5, NAVY, 8, 3),
    }
    for style_name, (size, color, before, after) in style_specs.items():
        style = document.styles[style_name]
        _set_style_font(style, "Calibri", "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Kicker" not in document.styles:
        kicker = document.styles.add_style("Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        kicker = document.styles["Kicker"]
    _set_style_font(kicker, "Calibri", "Microsoft YaHei")
    kicker.font.size = Pt(8)
    kicker.font.bold = True
    kicker.font.color.rgb = RGBColor.from_string(TEAL)
    kicker.paragraph_format.space_after = Pt(7)
    kicker.paragraph_format.keep_with_next = True

    _populate_running_header(section.header)
    _populate_running_header(section.even_page_header)
    _populate_running_footer(section.footer)
    _populate_running_footer(section.even_page_footer)


def _add_text(
    document: Document,
    text: str,
    *,
    style: str | None = None,
    bold_lead: str | None = None,
    color: str | None = None,
    size: float | None = None,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
) -> Any:
    """Add a styled paragraph with an optional bold lead-in."""
    paragraph = document.add_paragraph(style=style)
    if alignment is not None:
        paragraph.alignment = alignment
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        _set_run_font(lead)
        lead.font.bold = True
        body = paragraph.add_run(text[len(bold_lead) :])
        _set_run_font(body)
        runs = [lead, body]
    else:
        run = paragraph.add_run(text)
        _set_run_font(run)
        runs = [run]
    for run in runs:
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)
    return paragraph


def _add_bullet(document: Document, text: str, *, color: str = INK) -> Any:
    """Add a compact CommonMark-like bullet paragraph to the memo."""
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    _set_run_font(run)
    run.font.size = Pt(10.3)
    run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def _add_callout(
    document: Document,
    title: str,
    body: str,
    *,
    fill: str = PALE_BLUE,
    accent: str = BLUE,
) -> Any:
    """Add a one-cell callout with an accent border."""
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.65)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_shading(cell, fill)
    _set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    _set_cell_border(
        cell,
        start={"val": "single", "sz": "22", "space": "0", "color": accent},
    )
    title_p = cell.paragraphs[0]
    title_p.paragraph_format.space_after = Pt(3)
    title_run = title_p.add_run(title)
    _set_run_font(title_run)
    title_run.font.bold = True
    title_run.font.size = Pt(11.5)
    title_run.font.color.rgb = RGBColor.from_string(accent)
    body_p = cell.add_paragraph()
    body_p.paragraph_format.space_after = Pt(0)
    body_p.paragraph_format.line_spacing = 1.08
    body_run = body_p.add_run(body)
    _set_run_font(body_run)
    body_run.font.size = Pt(10)
    body_run.font.color.rgb = RGBColor.from_string(INK)
    _set_keep_table_rows(table)
    return table


def _add_metric_cards(document: Document, cards: list[tuple[str, str, str]]) -> Any:
    """Add a row of three compact metric cards."""
    table = document.add_table(rows=1, cols=len(cards))
    table.autofit = False
    widths = [Inches(2.15)] * len(cards)
    for idx, (label, value, note) in enumerate(cards):
        cell = table.cell(0, idx)
        cell.width = widths[idx]
        _set_cell_shading(cell, MIST)
        _set_cell_margins(cell, top=130, start=140, bottom=130, end=140)
        _set_cell_border(
            cell,
            bottom={"val": "single", "sz": "14", "space": "0", "color": TEAL},
        )
        p_label = cell.paragraphs[0]
        p_label.paragraph_format.space_after = Pt(2)
        label_run = p_label.add_run(label)
        _set_run_font(label_run)
        label_run.font.size = Pt(8)
        label_run.font.bold = True
        label_run.font.color.rgb = RGBColor.from_string(SLATE)
        p_value = cell.add_paragraph()
        p_value.paragraph_format.space_after = Pt(1)
        value_run = p_value.add_run(value)
        _set_run_font(value_run)
        value_run.font.size = Pt(18)
        value_run.font.bold = True
        value_run.font.color.rgb = RGBColor.from_string(NAVY)
        p_note = cell.add_paragraph()
        p_note.paragraph_format.space_after = Pt(0)
        note_run = p_note.add_run(note)
        _set_run_font(note_run)
        note_run.font.size = Pt(7.8)
        note_run.font.color.rgb = RGBColor.from_string(SLATE)
    _set_keep_table_rows(table)
    return table


def _add_table(
    document: Document,
    headers: list[str],
    rows: Iterable[Iterable[str]],
    *,
    widths: list[float] | None = None,
    font_size: float = 8.7,
) -> Any:
    """Add a compact, consistently styled research table."""
    data_rows = [list(row) for row in rows]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    _set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        _set_cell_shading(cell, NAVY)
        _set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        _set_run_font(run)
        run.font.size = Pt(font_size)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, values in enumerate(data_rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            _set_cell_shading(cell, WHITE if row_index % 2 == 0 else MIST)
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            _set_run_font(run)
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(INK)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    _set_keep_table_rows(table)
    return table


def _add_source(document: Document, text: str) -> Any:
    """Add a compact source/method note below a figure or table."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(f"来源与口径：{text}")
    _set_run_font(run)
    run.font.size = Pt(7.5)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(SLATE)
    return paragraph


def _add_picture(document: Document, path: Path, *, width: float = 6.55) -> Any:
    """Insert a report chart centered on the page."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    return paragraph


def _new_page(document: Document) -> None:
    """Start a new page without creating a new section."""
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _draw_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    """Draw a standard title block on a chart canvas."""
    draw.text((105, 60), title, fill=PIL_COLORS["navy"], font=_font(48, bold=True))
    draw.text((105, 124), subtitle, fill=PIL_COLORS["slate"], font=_font(26))


def _draw_zero_line(draw: ImageDraw.ImageDraw, x: int, y0: int, y1: int) -> None:
    """Draw a consistent zero-reference line."""
    draw.line((x, y0, x, y1), fill=PIL_COLORS["slate"], width=3)


def _save_chart(image: Image.Image, name: str) -> Path:
    """Save a report chart as an optimized RGB PNG."""
    path = FIGURE_DIR / name
    image.convert("RGB").save(path, format="PNG", optimize=True)
    return path


def _chart_ic(representation_ic: pd.DataFrame) -> Path:
    """Create a confidence-interval chart for the four representations."""
    labels = {
        "gm_level": "Level",
        "gm_delta_yoy": "Delta",
        "gm_innovation_raw": "Innovation (raw)",
        "gm_innovation_std": "Innovation (std)",
    }
    order = ["gm_level", "gm_delta_yoy", "gm_innovation_raw", "gm_innovation_std"]
    rows = representation_ic.set_index("feature").loc[order]
    image = Image.new("RGB", (1800, 970), PIL_COLORS["white"])
    draw = ImageDraw.Draw(image)
    _draw_title(
        draw,
        "图 1｜同一样本上的 Level / Delta / Innovation",
        "Spearman IC 与发行人聚类 bootstrap 95% 置信区间",
    )
    left, right, top, bottom = 470, 1660, 245, 805
    xmin, xmax = -0.55, 0.45

    def x_pos(value: float) -> int:
        return int(left + (value - xmin) / (xmax - xmin) * (right - left))

    for tick in [-0.5, -0.25, 0.0, 0.25]:
        x = x_pos(tick)
        draw.line((x, top, x, bottom), fill=PIL_COLORS["grid"], width=2)
        text = f"{tick:+.2f}"
        draw.text((x - 38, bottom + 26), text, fill=PIL_COLORS["slate"], font=_font(22))
    _draw_zero_line(draw, x_pos(0.0), top, bottom)
    y_positions = [310, 440, 570, 700]
    for y, (feature, row) in zip(y_positions, rows.iterrows(), strict=True):
        label = labels[feature]
        draw.text((105, y - 25), label, fill=PIL_COLORS["ink"], font=_font(29, bold=True))
        low, point, high = float(row.ci_low), float(row.spearman_ic), float(row.ci_high)
        draw.line((x_pos(low), y, x_pos(high), y), fill=PIL_COLORS["blue"], width=9)
        draw.line((x_pos(low), y - 17, x_pos(low), y + 17), fill=PIL_COLORS["blue"], width=5)
        draw.line((x_pos(high), y - 17, x_pos(high), y + 17), fill=PIL_COLORS["blue"], width=5)
        marker = PIL_COLORS["teal"] if "Innovation" in label else PIL_COLORS["navy"]
        draw.ellipse((x_pos(point) - 14, y - 14, x_pos(point) + 14, y + 14), fill=marker)
        annotation = f"IC {point:+.3f}  [{low:+.3f}, {high:+.3f}]"
        draw.text((1180, y - 24), annotation, fill=PIL_COLORS["ink"], font=_font(23))
    draw.text(
        (105, 885),
        "四条置信区间全部跨过 0；Innovation 的方向略正，但统计证据不足。",
        fill=PIL_COLORS["red"],
        font=_font(26, bold=True),
    )
    return _save_chart(image, "figure_1_representation_ic.png")


def _chart_horizons(robustness: pd.DataFrame) -> Path:
    """Create a primary-versus-secondary horizon comparison."""
    selected = robustness[robustness.diagnostic_id.isin(["horizon_1", "horizon_5", "full_primary"])]
    indexed = selected.set_index("horizon_sessions").loc[[1, 5, 20]]
    image = Image.new("RGB", (1800, 970), PIL_COLORS["white"])
    draw = ImageDraw.Draw(image)
    _draw_title(
        draw,
        "图 2｜事件反应是否集中在更短窗口？",
        "1 日与 5 日为预注册次要窗口；20 日为锁定主检验窗口",
    )
    left, right, top, bottom = 210, 1670, 260, 800
    ymin, ymax = -0.10, 0.25

    def y_pos(value: float) -> int:
        return int(bottom - (value - ymin) / (ymax - ymin) * (bottom - top))

    for tick in [-0.10, 0.0, 0.10, 0.20]:
        y = y_pos(tick)
        draw.line((left, y, right, y), fill=PIL_COLORS["grid"], width=2)
        draw.text((105, y - 15), f"{tick:+.2f}", fill=PIL_COLORS["slate"], font=_font(22))
    draw.line((left, y_pos(0.0), right, y_pos(0.0)), fill=PIL_COLORS["slate"], width=3)
    centers = [470, 930, 1390]
    bar_width = 120
    for center, (horizon, row) in zip(centers, indexed.iterrows(), strict=True):
        for offset, value, color in [
            (-75, float(row.delta_ic), PIL_COLORS["navy"]),
            (75, float(row.innovation_ic), PIL_COLORS["teal"]),
        ]:
            x0, x1 = center + offset - bar_width // 2, center + offset + bar_width // 2
            y_zero, y_value = y_pos(0.0), y_pos(value)
            draw.rectangle((x0, min(y_zero, y_value), x1, max(y_zero, y_value)), fill=color)
            draw.text(
                (x0 - 2, y_value - 42 if value >= 0 else y_value + 10),
                f"{value:+.3f}",
                fill=color,
                font=_font(22, bold=True),
            )
        draw.text(
            (center - 70, bottom + 35),
            f"{horizon} 日",
            fill=PIL_COLORS["ink"],
            font=_font(28, bold=True),
        )
    draw.rectangle((1160, 178, 1195, 213), fill=PIL_COLORS["navy"])
    draw.text((1210, 176), "Delta", fill=PIL_COLORS["ink"], font=_font(23))
    draw.rectangle((1370, 178, 1405, 213), fill=PIL_COLORS["teal"])
    draw.text((1420, 176), "Innovation", fill=PIL_COLORS["ink"], font=_font(23))
    draw.text(
        (105, 885),
        "5 日窗口有探索性正向迹象，但经多重检验校正后仍未显著，不能事后替换主窗口。",
        fill=PIL_COLORS["gold"],
        font=_font(25, bold=True),
    )
    return _save_chart(image, "figure_2_horizon_comparison.png")


def _chart_quantiles(quantiles: pd.DataFrame) -> Path:
    """Create the standardized Innovation quantile-return chart."""
    rows = quantiles[quantiles.feature == "gm_innovation_std"].set_index("quantile")
    ordered = rows.loc[["Low", "Middle", "High"]]
    image = Image.new("RGB", (1800, 970), PIL_COLORS["white"])
    draw = ImageDraw.Draw(image)
    _draw_title(
        draw,
        "图 3｜Innovation 分位数组合并不单调",
        "20 个交易日市场超额收益均值；Low / Middle / High 为同一样本三分位",
    )
    left, right, top, bottom = 210, 1660, 250, 800
    ymin, ymax = -0.04, 0.02

    def y_pos(value: float) -> int:
        return int(bottom - (value - ymin) / (ymax - ymin) * (bottom - top))

    for tick in [-0.04, -0.02, 0.0, 0.02]:
        y = y_pos(tick)
        draw.line((left, y, right, y), fill=PIL_COLORS["grid"], width=2)
        draw.text((95, y - 14), _pct(tick, 0), fill=PIL_COLORS["slate"], font=_font(22))
    draw.line((left, y_pos(0.0), right, y_pos(0.0)), fill=PIL_COLORS["slate"], width=3)
    centers = [470, 930, 1390]
    labels = ["Low", "Middle", "High"]
    colors = [PIL_COLORS["red"], PIL_COLORS["blue"], PIL_COLORS["teal"]]
    for center, label, color in zip(centers, labels, colors, strict=True):
        value = float(ordered.loc[label].mean_market_excess_return)
        y_zero, y_value = y_pos(0.0), y_pos(value)
        draw.rectangle(
            (center - 145, min(y_zero, y_value), center + 145, max(y_zero, y_value)), fill=color
        )
        draw.text(
            (center - 88, y_value - 48 if value >= 0 else y_value + 15),
            _pct(value),
            fill=color,
            font=_font(28, bold=True),
        )
        n_events = int(ordered.loc[label].events)
        draw.text(
            (center - 80, bottom + 35), label, fill=PIL_COLORS["ink"], font=_font(28, bold=True)
        )
        draw.text(
            (center - 54, bottom + 82), f"n={n_events}", fill=PIL_COLORS["slate"], font=_font(21)
        )
    draw.text(
        (105, 905),
        "Middle (+0.92%) 高于 High (+0.42%)：当前样本不能支持清晰的线性排序交易规则。",
        fill=PIL_COLORS["red"],
        font=_font(25, bold=True),
    )
    return _save_chart(image, "figure_3_innovation_quantiles.png")


def _chart_survivorship(survivorship: pd.DataFrame) -> Path:
    """Create the survivorship-bias audit chart."""
    order = ["all_events", "survivors_only_biased", "later_delisted_retained"]
    rows = survivorship.set_index("sample_id").loc[order]
    image = Image.new("RGB", (1800, 970), PIL_COLORS["white"])
    draw = ImageDraw.Draw(image)
    _draw_title(
        draw,
        "图 4｜删除后来退市公司，会把结果人为抬高",
        "20 个交易日市场超额收益均值；退市公司按事件时点真实保留",
    )
    left, right, top, bottom = 210, 1660, 250, 790
    ymin, ymax = -0.16, 0.04

    def y_pos(value: float) -> int:
        return int(bottom - (value - ymin) / (ymax - ymin) * (bottom - top))

    for tick in [-0.15, -0.10, -0.05, 0.0]:
        y = y_pos(tick)
        draw.line((left, y, right, y), fill=PIL_COLORS["grid"], width=2)
        draw.text((80, y - 14), _pct(tick, 0), fill=PIL_COLORS["slate"], font=_font(22))
    draw.line((left, y_pos(0.0), right, y_pos(0.0)), fill=PIL_COLORS["slate"], width=3)
    centers = [470, 930, 1390]
    labels = ["全部事件", "只留存活者\n（偏差样本）", "后来退市者\n（应保留）"]
    colors = [PIL_COLORS["blue"], PIL_COLORS["gold"], PIL_COLORS["red"]]
    for center, label, color, (_, row) in zip(
        centers, labels, colors, rows.iterrows(), strict=True
    ):
        value = float(row.mean_market_excess_return)
        y_zero, y_value = y_pos(0.0), y_pos(value)
        draw.rectangle(
            (center - 145, min(y_zero, y_value), center + 145, max(y_zero, y_value)), fill=color
        )
        draw.text(
            (center - 92, y_value - 46 if value >= 0 else y_value + 14),
            _pct(value),
            fill=color,
            font=_font(28, bold=True),
        )
        line_1, *line_2 = label.split("\n")
        draw.text(
            (center - 105, bottom + 32), line_1, fill=PIL_COLORS["ink"], font=_font(25, bold=True)
        )
        if line_2:
            draw.text(
                (center - 105, bottom + 70), line_2[0], fill=PIL_COLORS["slate"], font=_font(21)
            )
        draw.text(
            (center - 80, bottom + 110),
            f"n={int(row.events)}",
            fill=PIL_COLORS["slate"],
            font=_font(20),
        )
    inflation = float(rows.loc["survivors_only_biased"].mean_market_excess_return) - float(
        rows.loc["all_events"].mean_market_excess_return
    )
    draw.rounded_rectangle((1035, 165, 1660, 232), radius=16, fill="#FFF4D6")
    draw.text(
        (1060, 178),
        f"幸存者样本抬高均值：+{inflation * 100:.2f} 个百分点",
        fill=PIL_COLORS["gold"],
        font=_font(25, bold=True),
    )
    return _save_chart(image, "figure_4_survivorship_bias.png")


def _chart_quarters(robustness: pd.DataFrame) -> Path:
    """Create a calendar-quarter stability chart for Innovation IC."""
    quarter_rows = robustness[robustness.diagnostic_id.str.startswith("calendar_quarter_")].copy()
    quarter_rows["quarter"] = quarter_rows.diagnostic_id.str.replace(
        "calendar_quarter_", "", regex=False
    )
    image = Image.new("RGB", (1800, 970), PIL_COLORS["white"])
    draw = ImageDraw.Draw(image)
    _draw_title(
        draw,
        "图 5｜Innovation 方向随季度明显翻转",
        "20 日 Innovation IC；每个柱仅包含 10-15 个事件，应视为稳定性诊断",
    )
    left, right, top, bottom = 210, 1660, 250, 800
    ymin, ymax = -0.50, 0.50

    def y_pos(value: float) -> int:
        return int(bottom - (value - ymin) / (ymax - ymin) * (bottom - top))

    for tick in [-0.5, -0.25, 0.0, 0.25, 0.5]:
        y = y_pos(tick)
        draw.line((left, y, right, y), fill=PIL_COLORS["grid"], width=2)
        draw.text((105, y - 14), f"{tick:+.2f}", fill=PIL_COLORS["slate"], font=_font(22))
    draw.line((left, y_pos(0.0), right, y_pos(0.0)), fill=PIL_COLORS["slate"], width=3)
    centers = [385, 735, 1085, 1435]
    for center, (_, row) in zip(centers, quarter_rows.iterrows(), strict=True):
        value = float(row.innovation_ic)
        color = PIL_COLORS["teal"] if value >= 0 else PIL_COLORS["red"]
        y_zero, y_value = y_pos(0.0), y_pos(value)
        draw.rectangle(
            (center - 120, min(y_zero, y_value), center + 120, max(y_zero, y_value)), fill=color
        )
        draw.text(
            (center - 70, y_value - 45 if value >= 0 else y_value + 12),
            f"{value:+.3f}",
            fill=color,
            font=_font(25, bold=True),
        )
        draw.text(
            (center - 72, bottom + 38),
            str(row.quarter),
            fill=PIL_COLORS["ink"],
            font=_font(24, bold=True),
        )
        draw.text(
            (center - 48, bottom + 77),
            f"n={int(row.events)}",
            fill=PIL_COLORS["slate"],
            font=_font(20),
        )
    draw.text(
        (105, 900),
        "2025Q2 为 -0.420，而相邻季度为正：当前证据具有明显状态依赖与小样本波动。",
        fill=PIL_COLORS["red"],
        font=_font(25, bold=True),
    )
    return _save_chart(image, "figure_5_quarter_stability.png")


def _chart_repair() -> Path:
    """Create a before-versus-after data-repair audit chart."""
    image = Image.new("RGB", (1800, 900), PIL_COLORS["white"])
    draw = ImageDraw.Draw(image)
    _draw_title(
        draw,
        "图 6｜补回缺失事件后，结果更保守",
        "OB→TEAD 代码迁移修复前后；主结论均为“尚未建立有效性”",
    )
    left, right, top, bottom = 250, 1650, 245, 710
    ymin, ymax = -0.05, 0.12

    def y_pos(value: float) -> int:
        return int(bottom - (value - ymin) / (ymax - ymin) * (bottom - top))

    for tick in [-0.05, 0.0, 0.05, 0.10]:
        y = y_pos(tick)
        draw.line((left, y, right, y), fill=PIL_COLORS["grid"], width=2)
        draw.text((120, y - 14), f"{tick:+.2f}", fill=PIL_COLORS["slate"], font=_font(22))
    draw.line((left, y_pos(0.0), right, y_pos(0.0)), fill=PIL_COLORS["slate"], width=3)
    centers = [650, 1250]
    points = [
        ("修复前\n48 个事件", 0.093793, PIL_COLORS["gold"]),
        ("修复后\n49 个事件", 0.044695, PIL_COLORS["teal"]),
    ]
    for center, (label, value, color) in zip(centers, points, strict=True):
        y_zero, y_value = y_pos(0.0), y_pos(value)
        draw.rectangle(
            (center - 165, min(y_zero, y_value), center + 165, max(y_zero, y_value)), fill=color
        )
        draw.text(
            (center - 88, y_value - 48), f"{value:+.3f}", fill=color, font=_font(30, bold=True)
        )
        first, second = label.split("\n")
        draw.text(
            (center - 75, bottom + 35), first, fill=PIL_COLORS["ink"], font=_font(27, bold=True)
        )
        draw.text((center - 92, bottom + 78), second, fill=PIL_COLORS["slate"], font=_font(22))
    draw.line((815, 340, 1080, 430), fill=PIL_COLORS["red"], width=7)
    draw.polygon([(1080, 430), (1035, 392), (1040, 448)], fill=PIL_COLORS["red"])
    draw.text((835, 285), "-0.049", fill=PIL_COLORS["red"], font=_font(28, bold=True))
    draw.text(
        (105, 835),
        "数据质量修复没有迎合正向结果；它恢复了真实可交易路径，并削弱了原先较乐观的估计。",
        fill=PIL_COLORS["navy"],
        font=_font(25, bold=True),
    )
    return _save_chart(image, "figure_6_transition_repair.png")


def _build_figures(data: dict[str, pd.DataFrame]) -> dict[str, Path]:
    """Build the complete chart pack from locked diagnostic tables."""
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    return {
        "ic": _chart_ic(data["representation_ic"]),
        "horizons": _chart_horizons(data["robustness"]),
        "quantiles": _chart_quantiles(data["quantiles"]),
        "survivorship": _chart_survivorship(data["survivorship"]),
        "quarters": _chart_quarters(data["robustness"]),
        "repair": _chart_repair(),
    }


def _add_cover(document: Document, report: dict[str, Any]) -> None:
    """Build the memo_masthead first page and locked executive summary."""
    _add_text(document, "BIAS-AWARE RESEARCH NOTE · LOCKED V1", style="Kicker")
    title = document.add_paragraph(style="Title")
    title_run = title.add_run("财报信号初步有效性报告")
    _set_run_font(title_run, east_asia="Microsoft YaHei")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle_run = subtitle.add_run(
        "Level vs Delta vs Innovation｜真实财报事件、点时数据与偏差审计"
    )
    _set_run_font(subtitle_run)

    meta = _add_table(
        document,
        ["研究状态", "样本", "主窗口", "数据截点"],
        [["预注册后冻结", "49 事件 / 15 发行人", "20 个交易日", "2026-08-14"]],
        widths=[1.58, 1.75, 1.58, 1.75],
        font_size=8.2,
    )
    meta.rows[0].height = Cm(0.58)

    document.add_paragraph().paragraph_format.space_after = Pt(1)
    _add_callout(
        document,
        "正式结论：尚未建立 Innovation 的有效性",
        (
            "标准化毛利率 Innovation 的 20 日 IC 为 +0.045，但 95% 置信区间为 "
            "[-0.316, +0.356]；相对 Delta 的优势同样跨过 0。当前证据支持继续扩样，"
            "不支持据此建立买卖规则。"
        ),
        fill=PALE_RED,
        accent=RED,
    )
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    _add_metric_cards(
        document,
        [
            ("INNOVATION IC", "+0.045", "95% CI -0.316 至 +0.356"),
            ("相对 DELTA", "+0.064", "95% CI -0.274 至 +0.374"),
            ("幸存者偏差", "+1.81 pp", "只留存活公司抬高均值"),
        ],
    )

    _add_text(document, "一页摘要", style="Heading 2")
    _add_bullet(
        document,
        "研究设计有效：同一完整样本比较 Level、Delta 与严格样本外 Innovation，并锁定可交易时间窗。",
    )
    _add_bullet(
        document, "统计结论克制：7 项预注册判据仅通过 4 项，主 IC 与相对 Delta 的增量证据均不显著。"
    )
    _add_bullet(
        document, "经济排序不干净：Innovation 中分位收益高于高分位，暂不能转化为单调多空规则。"
    )
    _add_bullet(
        document, "偏差发现明确：剔除后来退市公司会把样本平均市场超额收益从 -0.57% 抬至 +1.24%。"
    )
    _add_bullet(
        document,
        (
            "数据审计有价值：补回 OB→TEAD 的真实价格路径后，Innovation IC 从 "
            "+0.094 降至 +0.045，结论仍未建立。"
        ),
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(
        f"版本：{report['dataset_version']}｜用途：阶段性研究沟通｜不构成投资建议或收益承诺"
    )
    _set_run_font(run)
    run.font.size = Pt(7.6)
    run.font.color.rgb = RGBColor.from_string(SLATE)


def _add_design_page(
    document: Document,
    prereg: dict[str, Any],
    ledger: dict[str, Any],
) -> None:
    """Add the pre-registered design and point-in-time sample definition."""
    _new_page(document)
    _add_text(document, "01｜研究问题与预注册设计", style="Heading 1")
    _add_callout(
        document,
        "主问题",
        (
            "在同一批 Company × Filing Event × Information Set 上，标准化样本外毛利率 "
            "Innovation 是否比毛利率同比变化 Delta 更能解释财报后 20 个交易日的市场超额收益？"
        ),
        fill=PALE_BLUE,
        accent=BLUE,
    )

    _add_text(document, "三种表示不是同一件事", style="Heading 2")
    _add_table(
        document,
        ["输入层", "定义", "回答的问题", "主要风险"],
        [
            ["Level", "本期毛利率水平", "公司现在处于什么状态？", "行业与商业模式差异"],
            ["Delta", "相对去年同期的变化", "状态改善还是恶化？", "季节性与基数效应"],
            [
                "Innovation",
                "实际值 - 事前样本外预测",
                "这次公告带来了多少新信息？",
                "预测模型与估计误差",
            ],
        ],
        widths=[1.0, 1.85, 2.1, 1.65],
    )
    _add_source(document, "预注册 estimand 与同样本比较契约；Innovation 不等同于分析师预期差。")

    _add_text(document, "锁定口径", style="Heading 2")
    _add_table(
        document,
        ["项目", "预先锁定的规则"],
        [
            ["原子单位", "公司 × 财报事件 × 当时可得信息集；三种表示使用同一完整样本。"],
            ["信号可得性", "仅使用原始披露值；预测与标准化误差均只使用目标事件之前的信息。"],
            ["入场", "申报接受时间之后的第二个未来正常交易日开盘价。"],
            ["退出", "第 20 个持有交易日收盘价；1 日与 5 日仅为次要窗口。"],
            ["基准", "同期 SPY 点时总收益；样本包含后来退市证券。"],
            ["不在范围", "不计算年化收益、夏普或策略 P&L；本轮是信号解释力检验。"],
        ],
        widths=[1.25, 5.35],
        font_size=8.5,
    )
    _add_source(
        document,
        (
            f"{prereg['dataset_version']}；{ledger['dataset_version']}。"
            "收益窗口使用未调整常规交易时段开盘/收盘，并显式处理拆股、现金股息与退市结局。"
        ),
    )

    _add_callout(
        document,
        "为什么这一步比直接回测更重要",
        "如果先看收益再换窗口、删异常公司或选择样本，研究会把数据偶然性伪装成可交易规律。这里先冻结问题与判据，再打开收益。",
        fill=PALE_TEAL,
        accent=TEAL,
    )


def _add_sample_page(document: Document, prereg: dict[str, Any], ledger: dict[str, Any]) -> None:
    """Add sample construction, lineage, and return-readiness diagnostics."""
    _new_page(document)
    _add_text(document, "02｜数据资产与样本形成", style="Heading 1")
    _add_text(
        document,
        "本轮不是从股票池里挑选“最好看”的事件，而是从预注册事件集合出发，逐层通过财报特征、身份映射、交易日历、公司行动、退市结局和基准收益门。",
    )
    _add_metric_cards(
        document,
        [
            ("候选事件", "58", "预注册时收益未打开"),
            ("完整特征+收益", "49", "超过门槛 47"),
            ("发行人", "15", "bootstrap 按发行人聚类"),
        ],
    )

    _add_text(document, "样本漏斗", style="Heading 2")
    _add_table(
        document,
        ["层级", "事件数", "覆盖率 / 状态", "控制目的"],
        [
            ["预注册候选事件", "58", "100%", "冻结选择集合，防止按收益挑样本"],
            [
                "完整 Level/Delta/Innovation",
                "49",
                f"{prereg['complete_feature_coverage']:.1%}",
                "确保同样本可比",
            ],
            ["20 日市场超额收益可用", "58", "全部候选事件", "确保价格路径与基准完整"],
            ["主检验完整事件", "49", "门槛 47；Gate=True", "同时具备特征与收益"],
            [
                "后来退市事件",
                str(ledger["later_delisted_events_retained"]),
                "显式保留",
                "避免幸存者偏差",
            ],
        ],
        widths=[1.65, 0.75, 1.55, 2.65],
    )
    _add_source(document, "预注册 event registry、最终事件收益账本与退市结果表。")

    _add_text(document, "点时可交易收益账本", style="Heading 2")
    _add_table(
        document,
        ["组件", "实现"],
        [
            ["公告时点", "SEC filing acceptance timestamp，不把 filed date 当作可得时点。"],
            ["证券身份", "CIK 到交易证券的历史映射；代码迁移跨段拼接必须通过证据门。"],
            ["公司收益", "未调整价格 + 拆股股数调整 + 现金股息保留，不把退市收益默认为 0。"],
            ["市场超额收益", "同一入场/退出交易日上的公司总收益减 SPY 总收益。"],
            ["不确定性", "10,000 次发行人聚类 bootstrap，避免把同公司多事件当作独立样本。"],
        ],
        widths=[1.35, 5.25],
    )
    _add_source(document, "事件收益账本 174 行 = 58 事件 × 3 个窗口；主分析仅使用 20 日完整样本。")

    _add_callout(
        document,
        "可审计性",
        (
            "每一层均保存版本号、manifest SHA-256、原始市场响应与拒绝原因；"
            "缺失或身份不确定不会被静默填补。"
        ),
        fill=PALE_TEAL,
        accent=TEAL,
    )


def _add_main_result_page(
    document: Document,
    figures: dict[str, Path],
    representation_ic: pd.DataFrame,
) -> None:
    """Add the main same-sample representation comparison."""
    _new_page(document)
    _add_text(document, "03｜主检验：Innovation 方向略正，但证据不足", style="Heading 1")
    _add_picture(document, figures["ic"])
    _add_source(
        document, "49 个完整事件、15 家发行人；Spearman IC；10,000 次发行人聚类 bootstrap 95% CI。"
    )

    labels = {
        "gm_level": "Level",
        "gm_delta_yoy": "Delta",
        "gm_innovation_raw": "Innovation raw",
        "gm_innovation_std": "Innovation std（主信号）",
    }
    rows = []
    for row in representation_ic.itertuples(index=False):
        rows.append(
            [
                labels[row.feature],
                str(row.events),
                f"{row.spearman_ic:+.3f}",
                f"[{row.ci_low:+.3f}, {row.ci_high:+.3f}]",
                "是" if row.ci_low <= 0 <= row.ci_high else "否",
            ]
        )
    _add_table(
        document,
        ["表示", "事件", "IC", "95% CI", "区间跨 0"],
        rows,
        widths=[2.0, 0.75, 0.85, 1.65, 1.0],
    )
    _add_source(document, "锁定输出 representation_ic.parquet；四种表示严格使用同一完整样本。")

    _add_callout(
        document,
        "怎样读这张图",
        (
            "点估计回答“样本里方向多大”，置信区间回答“换一批类似发行人后方向"
            "是否稳定”。Innovation 的点在 0 右侧，但区间很宽并跨过 0，因此不能排除"
            "无效或反向。"
        ),
        fill=PALE_GOLD,
        accent=GOLD,
    )


def _add_incremental_page(
    document: Document,
    figures: dict[str, Path],
    incremental: pd.DataFrame,
    multiplicity: pd.DataFrame,
) -> None:
    """Add incremental-value and horizon analyses."""
    _new_page(document)
    _add_text(document, "04｜增量价值与时间窗口", style="Heading 1")
    _add_text(
        document,
        (
            "Innovation 必须不仅“自己略正”，还要证明它比简单的 Delta 多提供了信息。"
            "两项增量检验的区间均跨过 0，bootstrap p 值约为 0.70。"
        ),
    )
    rows = []
    name_map = {
        "innovation_minus_delta_spearman_ic": "IC(Innovation) - IC(Delta)",
        "incremental_innovation_rank_beta_given_delta": "控制 Delta 后的 Innovation rank beta",
    }
    for row in incremental.itertuples(index=False):
        rows.append(
            [
                name_map[row.statistic],
                f"{row.point_estimate:+.3f}",
                f"[{row.ci_low:+.3f}, {row.ci_high:+.3f}]",
                f"{row.bootstrap_p_value:.3f}",
            ]
        )
    _add_table(
        document,
        ["增量检验", "点估计", "95% CI", "Bootstrap p"],
        rows,
        widths=[2.95, 1.05, 1.65, 1.0],
    )
    _add_source(document, "同一 49 事件样本；rank 回归与 IC 差异均按发行人聚类 bootstrap。")

    _add_picture(document, figures["horizons"])
    _add_source(document, "1 日与 5 日是预注册次要窗口；20 日为唯一确认性主窗口。")

    p_map = {int(row.horizon_sessions): row for row in multiplicity.itertuples(index=False)}
    five_day = p_map[5]
    one_day = p_map[1]
    _add_callout(
        document,
        "5 日信号值得跟踪，但不能升级为结论",
        (
            f"5 日 Innovation-Delta IC 差为 +0.219，原始 p={five_day.bootstrap_p_value:.3f}，"
            f"BH 校正后 p={five_day.bh_adjusted_p_value:.3f}；1 日差为 -0.054，"
            f"BH p={one_day.bh_adjusted_p_value:.3f}。下一批独立样本可把 5 日窗口"
            "预先锁定为复核假设。"
        ),
        fill=PALE_GOLD,
        accent=GOLD,
    )


def _add_tradeability_page(
    document: Document,
    figures: dict[str, Path],
    robustness: pd.DataFrame,
) -> None:
    """Add quantile monotonicity and stability diagnostics."""
    _new_page(document)
    _add_text(document, "05｜从统计相关到可交易排序：目前还差一步", style="Heading 1")
    _add_picture(document, figures["quantiles"])
    _add_source(document, "标准化 Innovation 三分位；20 日市场超额收益均值；每组 16-17 个事件。")

    _add_text(document, "为什么“IC 略正”仍不足以交易", style="Heading 2")
    _add_bullet(document, "高分位平均超额收益 +0.42%，但中分位为 +0.92%；排序不是单调上升。")
    _add_bullet(document, "低分位为 -2.89%，主要差异可能来自负尾部，而不是高 Innovation 持续跑赢。")
    _add_bullet(document, "49 个事件无法可靠估计换手、成本、容量、行业中性和组合回撤。")
    _add_bullet(document, "因此本轮报告不计算年化收益、夏普或策略净值，也不生成买卖建议。")

    _add_text(document, "跨季度稳定性", style="Heading 2")
    _add_picture(document, figures["quarters"])
    _add_source(document, "按财报接受时间划分季度；每季 10-15 个事件，仅用于检查状态依赖。")
    _add_callout(
        document,
        "稳定性风险",
        (
            "Innovation IC 在 2025Q2 为 -0.420，而其他三季为正。这个翻转可能来自"
            "宏观状态、行业构成或小样本噪声，必须通过扩样与时间外复核来区分。"
        ),
        fill=PALE_RED,
        accent=RED,
    )


def _add_bias_page(
    document: Document,
    figures: dict[str, Path],
    survivorship: pd.DataFrame,
) -> None:
    """Add the survivorship-bias result and research interpretation."""
    _new_page(document)
    _add_text(document, "06｜最明确的阶段性发现：幸存者偏差很大", style="Heading 1")
    _add_picture(document, figures["survivorship"])
    _add_source(document, "同一 49 事件完整样本；后来退市证券按事件发生时的真实可投资状态保留。")

    indexed = survivorship.set_index("sample_id")
    all_mean = float(indexed.loc["all_events"].mean_market_excess_return)
    survivor_mean = float(indexed.loc["survivors_only_biased"].mean_market_excess_return)
    inflation = survivor_mean - all_mean
    _add_callout(
        document,
        f"删除 6 个后来退市事件，样本均值被抬高 {inflation * 100:.2f} 个百分点",
        (
            f"完整样本平均市场超额收益为 {_pct(all_mean)}；只保留回看时仍存活的公司后变为 "
            f"{_pct(survivor_mean)}。这不是信号本身的盈利证明，而是样本选择会改变结论的直接证据。"
        ),
        fill=PALE_GOLD,
        accent=GOLD,
    )

    _add_text(document, "这一发现对项目有什么价值", style="Heading 2")
    _add_table(
        document,
        ["常见做法", "看似合理", "实际后果", "本框架的处理"],
        [
            [
                "用今天的股票池回测过去",
                "代码与数据更容易获取",
                "丢掉失败公司，抬高历史收益",
                "按事件时点保留后来退市证券",
            ],
            [
                "退市后价格缺失填 0 收益",
                "避免报错",
                "系统性低估尾部损失",
                "显式追踪退市结局，不静默插值",
            ],
            [
                "只报告完整价格的公司",
                "样本看起来更整齐",
                "可得性决定样本，产生选择偏差",
                "价格门与拒绝原因可审计",
            ],
        ],
        widths=[1.55, 1.45, 1.75, 1.85],
        font_size=8.2,
    )
    _add_source(document, "偏差分析是产品核心；它可独立于信号最终是否有效而产生研究价值。")


def _add_repair_page(
    document: Document,
    figures: dict[str, Path],
    repair: dict[str, Any],
) -> None:
    """Add the OB-to-TEAD ticker-transition audit and impact."""
    _new_page(document)
    _add_text(document, "07｜数据修复审计：OB→TEAD 代码迁移", style="Heading 1")
    _add_text(
        document,
        (
            "一项已锁定事件的 20 日持有期跨过证券代码迁移。若只按旧 ticker 抓价，"
            "退出路径会被错误截断；本轮没有删除该事件，而是建立了独立的身份与价格证据门。"
        ),
    )
    _add_picture(document, figures["repair"])
    _add_source(document, "修复前为 48 个完整事件；修复后为 49 个。主结论在两版中均未建立。")

    _add_table(
        document,
        ["审计项", "证据", "结果"],
        [
            ["身份连续", f"同 CIK {repair['cik']} / security_id {repair['security_id']}", "通过"],
            ["交易日连续", f"2025-06-09 OB → {repair['transition_date']} TEAD", "通过"],
            [
                "边界价格",
                f"旧收盘 $2.74；新开盘 $2.79；隔夜 {_pct(repair['overnight_return'])}",
                "通过",
            ],
            ["价格补充", f"TEAD 未调整日线 {repair['price_addition_rows']} 行", "通过"],
            ["公司行动", f"补充 {repair['action_addition_rows']} 行；证据检查 6/6", "通过"],
            ["选择独立", "修复计划不使用事件收益；信号值未改变", "通过"],
        ],
        widths=[1.3, 4.25, 0.9],
        font_size=8.3,
    )
    _add_source(document, f"{repair['dataset_version']}；原始响应、hash 与拒绝原因均保留。")

    _add_callout(
        document,
        "修复为何提高可信度",
        (
            "补回负向事件后 Innovation IC 从 +0.094 降至 +0.045。研究管线接受了"
            "更不利的结果，说明数据质量规则先于结论，而不是为了得到正收益。"
        ),
        fill=PALE_TEAL,
        accent=TEAL,
    )


def _add_verdict_page(
    document: Document,
    verdict_checks: pd.DataFrame,
    robustness: pd.DataFrame,
) -> None:
    """Add the locked verdict checklist and honest interpretation."""
    _new_page(document)
    _add_text(document, "08｜正式判定与结论边界", style="Heading 1")
    _add_callout(
        document,
        "效力尚未建立（effectiveness_not_established）",
        (
            "7 项预注册判据通过 4 项。通过的是数据门、增量 beta 点估计、留一发行人"
            "方向比例和单一发行人影响上限；失败的是两项置信区间与 ASNS 尾部稳健性。"
        ),
        fill=PALE_RED,
        accent=RED,
    )

    check_names = {
        "data_quality_and_coverage": "数据质量与覆盖门",
        "innovation_ic_positive_and_ci_excludes_zero": "Innovation IC 为正且 CI 排除 0",
        "innovation_minus_delta_ci_excludes_zero": "Innovation-Delta 差异 CI 排除 0",
        "incremental_innovation_beta_positive": "控制 Delta 后 Innovation beta 为正",
        "leave_one_issuer_out_positive_sign_share": "留一发行人正向比例 ≥ 80%",
        "exclude_ASNS_innovation_sign_positive": "排除 ASNS 后 Innovation 仍为正",
        "maximum_single_issuer_absolute_effect_share": "单一发行人绝对影响占比 ≤ 25%",
    }
    rows = []
    for row in verdict_checks.itertuples(index=False):
        rows.append(
            [
                check_names[row.check_id],
                "通过" if row.passed else "未通过",
                f"{row.observed:+.3f}" if abs(float(row.observed)) < 1 else f"{row.observed:.0f}",
                f"{row.threshold:+.3f}"
                if abs(float(row.threshold)) < 1
                else f"{row.threshold:.0f}",
            ]
        )
    table = _add_table(
        document,
        ["预注册判据", "状态", "观察值", "阈值"],
        rows,
        widths=[3.8, 0.9, 0.9, 0.9],
        font_size=8.2,
    )
    for row_idx, source_row in enumerate(verdict_checks.itertuples(index=False), start=1):
        status_cell = table.cell(row_idx, 1)
        _set_cell_shading(status_cell, PALE_TEAL if source_row.passed else PALE_RED)
    _add_source(document, "判据在查看收益前锁定；不能因为 5 日窗口较好或点估计为正而改写。")

    exclude_asns = robustness[robustness.diagnostic_id == "exclude_ASNS_tail_case"].iloc[0]
    _add_text(document, "当前能说什么 / 不能说什么", style="Heading 2")
    _add_table(
        document,
        ["可以说", "不能说"],
        [
            [
                "Innovation 的方向在全样本略正，值得扩样复核。",
                "Innovation 已经是有效或可盈利的财报因子。",
            ],
            ["5 日反应可能比 20 日更集中，是下一轮假设。", "看到 5 日较好后把主窗口改成 5 日。"],
            ["幸存者偏差可显著扭曲样本均值。", "完整样本的负均值说明应做空整个股票池。"],
            [
                f"排除 ASNS 后 Innovation IC={exclude_asns.innovation_ic:+.3f}，尾部较脆弱。",
                "当前 49 个事件足以训练实盘买卖模型。",
            ],
        ],
        widths=[3.25, 3.25],
        font_size=8.5,
    )
    _add_source(document, "本报告是信号解释力研究，不是策略业绩或投资建议。")


def _add_next_steps_page(document: Document, report: dict[str, Any]) -> None:
    """Add the next-round research plan and immutable provenance appendix."""
    _new_page(document)
    _add_text(document, "09｜下一轮：增加独立证据，而不是调参", style="Heading 1")
    _add_text(
        document,
        "现阶段的最高边际价值是扩展独立发行人、行业和市场状态。模型复杂度不是主要瓶颈；样本量、点时预期和市场微观数据才是。",
    )
    _add_table(
        document,
        ["优先级", "工作包", "具体动作", "成功标准"],
        [
            [
                "P0",
                "独立扩样",
                "新增发行人和退市/清退企业；不按本轮收益筛选",
                "完整事件数显著扩大，行业分布更均衡",
            ],
            [
                "P0",
                "5 日外样本复核",
                "在新样本上事前锁定 5 日窗口，不改本轮 20 日主结论",
                "方向、CI 与分位数单调性同时改善",
            ],
            [
                "P1",
                "预期层升级",
                "加入分析师一致预期、公司指引与文本语义 surprise",
                "Innovation 不再主要是 Delta 的重表达",
            ],
            [
                "P1",
                "交易现实",
                "加入流动性、点差、成交量、隔夜跳空和成本模型",
                "从解释力过渡到可交易净收益",
            ],
            [
                "P2",
                "状态条件",
                "加入行业、利率、通胀、波动率和市场 regime",
                "解释季度方向翻转并完成外样本验证",
            ],
            [
                "P3",
                "模型复杂度",
                "先用线性/树模型；样本和标签成熟后再评估深度学习",
                "复杂模型稳定胜过透明基线才升级",
            ],
        ],
        widths=[0.55, 1.1, 3.15, 1.7],
        font_size=7.9,
    )
    _add_source(document, "下一轮保持数据规则与判定先于结果，避免把调参写成研究发现。")

    _add_text(document, "对实盘量化的含义", style="Heading 2")
    _add_bullet(
        document,
        "这套系统首先承担“数据与偏差防火墙”：确保训练标签、股票池、事件时点和收益窗口没有偷看未来。",
    )
    _add_bullet(
        document,
        "未来买卖模型需要财报之外的价格/成交量、流动性、预期、文本、宏观与风险暴露；财报信号只是事件层输入。",
    )
    _add_bullet(
        document,
        "在样本规模和实盘成本没有成熟前，透明基线通常比深度学习更可诊断，也更容易发现偏差。",
    )

    _add_text(document, "版本与可复现性", style="Heading 2")
    _add_table(
        document,
        ["冻结对象", "版本 / SHA-256"],
        [
            ["预注册", "preliminary-effectiveness-preregistration-v1-97df4b07e54937e8"],
            ["事件收益账本", "longitudinal-event-return-ledger-v1-295f4ff0fa9a202e"],
            ["最终诊断", str(report["dataset_version"])],
            ["预注册 manifest", str(report["preregistration_manifest_sha256"])],
            ["收益账本 manifest", str(report["ledger_manifest_sha256"])],
        ],
        widths=[1.65, 4.85],
        font_size=7.4,
    )
    _add_source(document, "所有核心表均为版本化不可变输出；原始响应、解析产物和 hash 分层保留。")


def _load_data() -> tuple[dict[str, Any], dict[str, pd.DataFrame]]:
    """Load the exact frozen artifacts used by the report."""
    reports = {
        "diagnostic": _read_json(DIAGNOSTIC_DIR / "effectiveness_diagnostic_report.json"),
        "prereg": _read_json(PREREG_DIR / "preregistration_report.json"),
        "ledger": _read_json(LEDGER_DIR / "longitudinal_return_report.json"),
        "repair": _read_json(REPAIR_DIR / "ticker_transition_repair_report.json"),
    }
    tables = {
        "representation_ic": pd.read_parquet(DIAGNOSTIC_DIR / "representation_ic.parquet"),
        "incremental": pd.read_parquet(DIAGNOSTIC_DIR / "incremental_test.parquet"),
        "robustness": pd.read_parquet(DIAGNOSTIC_DIR / "robustness_diagnostics.parquet"),
        "quantiles": pd.read_parquet(DIAGNOSTIC_DIR / "quantile_returns.parquet"),
        "multiplicity": pd.read_parquet(DIAGNOSTIC_DIR / "secondary_multiplicity.parquet"),
        "survivorship": pd.read_parquet(DIAGNOSTIC_DIR / "survivorship_diagnostic.parquet"),
        "verdict_checks": pd.read_parquet(DIAGNOSTIC_DIR / "verdict_checks.parquet"),
    }
    return reports, tables


def build_report() -> Path:
    """Build and save the locked Chinese preliminary effectiveness memo."""
    reports, tables = _load_data()
    if reports["diagnostic"]["verdict"] != "effectiveness_not_established":
        raise ValueError("The report builder only supports the locked v1 verdict.")
    if int(reports["diagnostic"]["primary_analysis_events"]) != 49:
        raise ValueError("Unexpected primary event count; refusing to relabel the report.")
    if not bool(reports["repair"]["transition_gate"]):
        raise ValueError("Ticker-transition repair gate did not pass.")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    figures = _build_figures(tables)

    document = Document()
    _configure_document(document)
    _add_cover(document, reports["diagnostic"])
    _add_design_page(document, reports["prereg"], reports["ledger"])
    _add_sample_page(document, reports["prereg"], reports["ledger"])
    _add_main_result_page(document, figures, tables["representation_ic"])
    _add_incremental_page(document, figures, tables["incremental"], tables["multiplicity"])
    _add_tradeability_page(document, figures, tables["robustness"])
    _add_bias_page(document, figures, tables["survivorship"])
    _add_repair_page(document, figures, reports["repair"])
    _add_verdict_page(document, tables["verdict_checks"], tables["robustness"])
    _add_next_steps_page(document, reports["diagnostic"])

    core_properties = document.core_properties
    core_properties.title = "财报信号初步有效性报告 v1"
    core_properties.subject = "Level vs Delta vs Innovation preliminary effectiveness"
    core_properties.author = "Filing-to-Signal Lab"
    core_properties.keywords = "point-in-time, SEC EDGAR, Innovation, survivorship bias"
    core_properties.comments = (
        "Generated from immutable diagnostic artifacts; educational research only."
    )
    document.save(DOCX_PATH)
    return DOCX_PATH


def main() -> None:
    """CLI entry point."""
    build_report()
    print("REPORT DOCX DONE")


if __name__ == "__main__":
    main()
