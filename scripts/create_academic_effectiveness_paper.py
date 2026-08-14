"""Create the academic-style working paper for the preliminary effectiveness study.

The paper is a reporting layer over immutable research artifacts. It does not
re-estimate signals, returns, confidence intervals, or the locked verdict.
"""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

from scripts import create_preliminary_effectiveness_report as brief

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "reports" / "academic_effectiveness_paper_v1"
WORK_DIR = ROOT / "tmp" / "academic_effectiveness_paper_v1"
FIGURE_DIR = WORK_DIR / "figures"
DOCX_PATH = OUTPUT_DIR / "财报事件信号的点时验证_阶段性工作论文_v1.docx"

DIAGNOSTIC_DIR = brief.DIAGNOSTIC_DIR
PREREG_DIR = brief.PREREG_DIR
LEDGER_DIR = brief.LEDGER_DIR
REPAIR_DIR = brief.REPAIR_DIR

NAVY = "0B2545"
BLUE = "2E74B5"
TEAL = "16817A"
GOLD = "9A6500"
RED = "9B1C1C"
INK = "172B2B"
SLATE = "5E6F78"
LIGHT = "F4F6F9"
GRID = "D8E0E5"
WHITE = "FFFFFF"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F4F2"
PALE_GOLD = "FFF4D6"
PALE_RED = "FBE9E7"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def _set_run_font(
    run: Any,
    *,
    latin: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
) -> None:
    """Set Latin and East Asian typefaces on a run."""
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _set_style_font(style: Any, latin: str, east_asia: str) -> None:
    """Set Latin and East Asian typefaces on a paragraph style."""
    style.font.name = latin
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _set_cell_margins(
    cell: Any,
    *,
    top: int = 80,
    bottom: int = 80,
    start: int = 120,
    end: int = 120,
) -> None:
    """Set table-cell margins in twips."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    values = {"top": top, "bottom": bottom, "start": start, "end": end}
    for name, value in values.items():
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_cell_shading(cell: Any, fill: str) -> None:
    """Apply a solid fill to a Word table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_border(cell: Any, **edges: dict[str, str]) -> None:
    """Set selected table-cell borders."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, values in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in values.items():
            edge.set(qn(f"w:{key}"), str(value))


def _set_table_geometry(table: Any, widths_dxa: list[int]) -> None:
    """Apply fixed 9360-DXA table width, indent, grid, and matching cell widths."""
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("Academic paper table widths must total 9360 DXA")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def _set_repeat_header(row: Any) -> None:
    """Mark a table row as a repeating header."""
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(header)


def _add_page_field(paragraph: Any) -> None:
    """Insert a live PAGE field into a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def _populate_header(header: Any) -> None:
    """Populate an odd/even running header."""
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("FILING-TO-SIGNAL LAB  |  WORKING PAPER NO. 1")
    _set_run_font(run)
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(SLATE)


def _populate_footer(footer: Any) -> None:
    """Populate an odd/even footer with status, disclaimer, and page number."""
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    _set_table_geometry(table, [7920, 1440])
    left = table.cell(0, 0).paragraphs[0]
    left_run = left.add_run("阶段性工作论文｜审阅稿｜不构成投资建议")
    _set_run_font(left_run)
    left_run.font.size = Pt(7.5)
    left_run.font.color.rgb = RGBColor.from_string(SLATE)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    page_run = right.add_run("P. ")
    _set_run_font(page_run)
    page_run.font.size = Pt(7.5)
    page_run.font.color.rgb = RGBColor.from_string(SLATE)
    _add_page_field(right)
    for cell in table.row_cells(0):
        _set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
        _set_cell_border(
            cell,
            top={"val": "single", "sz": "6", "space": "0", "color": GRID},
        )


def _configure_document(document: Document) -> None:
    """Apply the narrative_proposal preset and the paper's named overrides."""
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri", "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = document.styles[name]
        _set_style_font(style, "Calibri", "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Paper Title" not in document.styles:
        title = document.styles.add_style("Paper Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        title = document.styles["Paper Title"]
    _set_style_font(title, "Calibri", "Microsoft YaHei")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)

    if "Paper Subtitle" not in document.styles:
        subtitle = document.styles.add_style("Paper Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = document.styles["Paper Subtitle"]
    _set_style_font(subtitle, "Calibri", "Microsoft YaHei")
    subtitle.font.size = Pt(13)
    subtitle.font.italic = True
    subtitle.font.color.rgb = RGBColor.from_string(SLATE)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)

    if "Equation" not in document.styles:
        equation = document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        equation = document.styles["Equation"]
    _set_style_font(equation, "Cambria Math", "Microsoft YaHei")
    equation.font.size = Pt(10.5)
    equation.font.color.rgb = RGBColor.from_string(NAVY)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(5)
    equation.paragraph_format.space_after = Pt(7)
    equation.paragraph_format.keep_together = True

    if "Source Note" not in document.styles:
        source = document.styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source = document.styles["Source Note"]
    _set_style_font(source, "Calibri", "Microsoft YaHei")
    source.font.size = Pt(8)
    source.font.italic = True
    source.font.color.rgb = RGBColor.from_string(SLATE)
    source.paragraph_format.space_before = Pt(4)
    source.paragraph_format.space_after = Pt(4)

    for style_name in ["List Bullet", "List Number"]:
        style = document.styles[style_name]
        _set_style_font(style, "Calibri", "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    _populate_header(section.header)
    _populate_header(section.even_page_header)
    _populate_footer(section.footer)
    _populate_footer(section.even_page_footer)


def _paragraph(
    document: Document,
    text: str,
    *,
    style: str | None = None,
    bold_lead: str | None = None,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
) -> Any:
    """Add a body paragraph with optional bold lead-in."""
    paragraph = document.add_paragraph(style=style)
    if alignment is not None:
        paragraph.alignment = alignment
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        _set_run_font(lead)
        lead.font.bold = True
        body = paragraph.add_run(text[len(bold_lead) :])
        _set_run_font(body)
    else:
        run = paragraph.add_run(text)
        _set_run_font(run)
    return paragraph


def _heading(document: Document, text: str, level: int = 1) -> Any:
    """Add a real Word heading."""
    return _paragraph(document, text, style=f"Heading {level}")


def _bullet(document: Document, text: str) -> Any:
    """Add a real Word bullet item."""
    return _paragraph(document, text, style="List Bullet")


def _numbered(document: Document, text: str) -> Any:
    """Add a real Word numbered-list item."""
    return _paragraph(document, text, style="List Number")


def _manual_numbered(document: Document, number: int, text: str) -> Any:
    """Add a numbered item whose sequence must restart independently."""
    paragraph = _paragraph(document, f"{number}.  {text}")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.194)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208
    return paragraph


def _equation(document: Document, expression: str) -> Any:
    """Add a centered, Cambria Math equation paragraph."""
    return _paragraph(document, expression, style="Equation")


def _source(document: Document, text: str) -> Any:
    """Add a source or methodological note using the preset citation rhythm."""
    return _paragraph(document, f"注：{text}", style="Source Note")


def _page_break(document: Document) -> None:
    """Start a new page."""
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _callout(
    document: Document,
    title: str,
    body: str,
    *,
    fill: str = PALE_BLUE,
    accent: str = BLUE,
) -> Any:
    """Add a restrained one-column academic note box."""
    table = document.add_table(rows=1, cols=1)
    _set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    _set_cell_shading(cell, fill)
    _set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    _set_cell_border(
        cell,
        start={"val": "single", "sz": "20", "space": "0", "color": accent},
    )
    title_p = cell.paragraphs[0]
    title_p.paragraph_format.space_after = Pt(3)
    title_run = title_p.add_run(title)
    _set_run_font(title_run)
    title_run.font.bold = True
    title_run.font.size = Pt(11.5)
    title_run.font.color.rgb = RGBColor.from_string(accent)
    body_p = cell.add_paragraph()
    body_p.paragraph_format.space_after = Pt(0)
    body_p.paragraph_format.line_spacing = 1.15
    body_run = body_p.add_run(body)
    _set_run_font(body_run)
    body_run.font.size = Pt(10.3)
    body_run.font.color.rgb = RGBColor.from_string(INK)
    return table


def _table(
    document: Document,
    title: str,
    headers: list[str],
    rows: Iterable[Iterable[str]],
    widths_dxa: list[int],
    *,
    font_size: float = 8.7,
) -> Any:
    """Add a titled fixed-geometry academic table."""
    caption = document.add_paragraph()
    caption.paragraph_format.space_before = Pt(7)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    caption_run = caption.add_run(title)
    _set_run_font(caption_run)
    caption_run.font.bold = True
    caption_run.font.size = Pt(10)
    caption_run.font.color.rgb = RGBColor.from_string(NAVY)

    table = document.add_table(rows=1, cols=len(headers))
    _set_table_geometry(table, widths_dxa)
    table.style = "Table Grid"
    _set_repeat_header(table.rows[0])
    for index, label in enumerate(headers):
        cell = table.cell(0, index)
        _set_cell_shading(cell, LIGHT)
        _set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index > 0 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        _set_run_font(run)
        run.font.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(NAVY)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            _set_cell_shading(cell, WHITE if row_index % 2 == 0 else "FAFBFC")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if col_index > 0 and len(str(value)) < 28
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(str(value))
            _set_run_font(run)
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(INK)
    return table


def _picture(document: Document, path: Path, *, width: float = 6.4) -> Any:
    """Insert a centered academic figure."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    return paragraph


def _retitle_chart(source: Path, destination: Path, title: str, subtitle: str) -> Path:
    """Create an academic-numbered copy of an existing locked-result chart."""
    image = Image.open(source).convert("RGB")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, image.width, 180), fill="#FFFFFF")
    draw.text((105, 60), title, fill="#0B2545", font=brief._font(46, bold=True))
    draw.text((105, 122), subtitle, fill="#5E6F78", font=brief._font(25))
    image.save(destination, format="PNG", optimize=True)
    return destination


def _concept_figure() -> Path:
    """Draw the State/Delta/Innovation conceptual decomposition."""
    image = Image.new("RGB", (1800, 940), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text(
        (90, 55),
        "图 1｜市场持续整合信息，决策依赖事件创新",
        fill="#0B2545",
        font=brief._font(45, bold=True),
    )
    draw.text(
        (90, 118),
        "State 描述公告前的条件状态；Innovation 描述新观察值相对事前条件预期的偏离",
        fill="#5E6F78",
        font=brief._font(25),
    )

    y = 455
    draw.line((150, y, 1650, y), fill="#2E74B5", width=8)
    for x, label in [(240, "历史财报"), (560, "行业信息"), (880, "价格状态"), (1260, "公告前 t-")]:
        draw.ellipse((x - 14, y - 14, x + 14, y + 14), fill="#2E74B5")
        draw.text((x - 60, y + 40), label, fill="#172B2B", font=brief._font(24, bold=True))
    draw.text(
        (140, 320),
        "连续积分：形成条件状态与市场预期",
        fill="#2E74B5",
        font=brief._font(30, bold=True),
    )

    event_x = 1450
    draw.line((event_x, 250, event_x, 680), fill="#9B1C1C", width=8)
    draw.polygon([(event_x, 225), (event_x - 18, 265), (event_x + 18, 265)], fill="#9B1C1C")
    draw.text((1360, 190), "公告事件 t", fill="#9B1C1C", font=brief._font(30, bold=True))
    draw.rounded_rectangle(
        (1080, 620, 1710, 820), radius=22, fill="#FBE9E7", outline="#9B1C1C", width=4
    )
    draw.text((1120, 655), "瞬时微分：Innovation", fill="#9B1C1C", font=brief._font(31, bold=True))
    draw.text((1120, 713), "实际披露 - 公告前条件预期", fill="#172B2B", font=brief._font(27))
    draw.text((1120, 760), "决定事件驱动交易的新增信息", fill="#172B2B", font=brief._font(24))

    draw.rounded_rectangle(
        (110, 620, 930, 820), radius=22, fill="#E8F4F2", outline="#16817A", width=4
    )
    draw.text((155, 655), "两个输入层", fill="#16817A", font=brief._font(31, bold=True))
    draw.text(
        (155, 715), "State Features：Level、行业、历史波动", fill="#172B2B", font=brief._font(25)
    )
    draw.text(
        (155, 763),
        "Innovation Features：Delta、预测误差、surprise",
        fill="#172B2B",
        font=brief._font(25),
    )
    path = FIGURE_DIR / "figure_1_concept.png"
    image.save(path, format="PNG", optimize=True)
    return path


def _pipeline_figure() -> Path:
    """Draw the immutable point-in-time research pipeline and sample gates."""
    image = Image.new("RGB", (1800, 980), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text(
        (90, 55),
        "图 2｜从 SEC 原始披露到冻结统计结论",
        fill="#0B2545",
        font=brief._font(45, bold=True),
    )
    draw.text(
        (90, 118),
        "每个箭头都携带时间戳、来源哈希与拒绝原因；收益在预注册之后才接入",
        fill="#5E6F78",
        font=brief._font(25),
    )
    boxes = [
        ("SEC EDGAR", "10-K / 10-Q\nXBRL + acceptance"),
        ("点时特征", "原始值、同比滞后\n行业滞后信息"),
        ("样本外预测", "扩展窗口 pooled OLS\n严格历史误差 σ"),
        ("事件收益账本", "第二未来交易日开盘\n20 日收盘 + SPY"),
        ("偏差与推断", "退市保留、cluster bootstrap\n预注册判据"),
    ]
    x_positions = [90, 430, 770, 1110, 1450]
    for index, ((title, body), x) in enumerate(zip(boxes, x_positions, strict=True)):
        fill = "#EAF2F8" if index < 3 else "#E8F4F2"
        outline = "#2E74B5" if index < 3 else "#16817A"
        draw.rounded_rectangle(
            (x, 275, x + 260, 545), radius=20, fill=fill, outline=outline, width=4
        )
        draw.text((x + 25, 315), title, fill=outline, font=brief._font(27, bold=True))
        for line_index, line in enumerate(body.split("\n")):
            draw.text((x + 25, 390 + line_index * 48), line, fill="#172B2B", font=brief._font(22))
        if index < len(boxes) - 1:
            draw.line((x + 270, 410, x + 325, 410), fill="#5E6F78", width=5)
            draw.polygon([(x + 325, 410), (x + 305, 395), (x + 305, 425)], fill="#5E6F78")

    draw.rounded_rectangle(
        (160, 690, 1640, 875), radius=22, fill="#F4F6F9", outline="#D8E0E5", width=3
    )
    metrics = [
        ("候选事件", "58"),
        ("完整特征 + 收益", "49"),
        ("发行人", "15"),
        ("10-Q / 10-K", "38 / 11"),
        ("后来退市事件", "6"),
    ]
    for index, (label, value) in enumerate(metrics):
        x = 205 + index * 285
        draw.text((x, 725), label, fill="#5E6F78", font=brief._font(20, bold=True))
        draw.text((x, 770), value, fill="#0B2545", font=brief._font(37, bold=True))
    path = FIGURE_DIR / "figure_2_pipeline.png"
    image.save(path, format="PNG", optimize=True)
    return path


def _build_figures(tables: dict[str, pd.DataFrame]) -> dict[str, Path]:
    """Build conceptual and renumbered empirical figures for the paper."""
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    locked = brief._build_figures(tables)
    return {
        "concept": _concept_figure(),
        "pipeline": _pipeline_figure(),
        "ic": _retitle_chart(
            locked["ic"],
            FIGURE_DIR / "figure_3_ic.png",
            "图 3｜Level、Delta 与 Innovation 的主样本 IC",
            "Spearman IC；发行人聚类 bootstrap 95% 置信区间",
        ),
        "horizons": _retitle_chart(
            locked["horizons"],
            FIGURE_DIR / "figure_4_horizons.png",
            "图 4｜不同事件窗口的 Innovation 与 Delta",
            "1 日与 5 日为预注册次要窗口；20 日为主窗口",
        ),
        "quantiles": _retitle_chart(
            locked["quantiles"],
            FIGURE_DIR / "figure_5_quantiles.png",
            "图 5｜标准化 Innovation 三分位收益",
            "20 日市场超额收益均值；每组 16-17 个事件",
        ),
        "survivorship": _retitle_chart(
            locked["survivorship"],
            FIGURE_DIR / "figure_7_survivorship.png",
            "图 7｜幸存者偏差对样本收益的影响",
            "保留后来退市证券，与回看时仅保留存活公司的样本比较",
        ),
        "quarters": _retitle_chart(
            locked["quarters"],
            FIGURE_DIR / "figure_6_quarters.png",
            "图 6｜Innovation IC 的季度稳定性",
            "每季度仅 10-15 个事件；用于诊断状态依赖而非独立发现",
        ),
        "repair": _retitle_chart(
            locked["repair"],
            FIGURE_DIR / "figure_8_repair.png",
            "图 8｜OB→TEAD 证券代码迁移修复的结果影响",
            "补回一项真实负向事件后，Innovation IC 更接近于零",
        ),
    }


def _cover(document: Document) -> None:
    """Build the editorial_cover first page."""
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(86)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("FILING-TO-SIGNAL LAB · WORKING PAPER NO. 1")
    _set_run_font(run)
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(GOLD)

    title = document.add_paragraph(style="Paper Title")
    title_run = title.add_run("财报事件信号的点时验证")
    _set_run_font(title_run)
    subtitle = document.add_paragraph(style="Paper Subtitle")
    subtitle_run = subtitle.add_run("Level、Delta 与 Innovation 的偏差感知研究")
    _set_run_font(subtitle_run)
    english = document.add_paragraph(style="Paper Subtitle")
    english_run = english.add_run(
        "A Bias-Aware Point-in-Time Test of State, Change, and Innovation Features"
    )
    _set_run_font(english_run)
    english_run.font.size = Pt(11.5)

    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.space_before = Pt(30)
    status.paragraph_format.space_after = Pt(78)
    status_run = status.add_run("阶段性工作论文 · 供专业审阅 · 2026 年 8 月 14 日")
    _set_run_font(status_run)
    status_run.font.size = Pt(10.5)
    status_run.font.bold = True
    status_run.font.color.rgb = RGBColor.from_string(TEAL)

    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(4)
    author_run = author.add_run("Filing-to-Signal Lab")
    _set_run_font(author_run)
    author_run.font.size = Pt(12)
    author_run.font.bold = True
    author_run.font.color.rgb = RGBColor.from_string(NAVY)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run("SEC EDGAR 原始披露 · Point-in-Time Fundamentals · Event Return Ledger")
    _set_run_font(meta_run)
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor.from_string(SLATE)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(50)
    note_run = note.add_run(
        "本稿报告一项预注册的技术性试点研究，不构成投资建议、收益承诺或已验证交易策略。"
    )
    _set_run_font(note_run)
    note_run.font.size = Pt(9)
    note_run.font.italic = True
    note_run.font.color.rgb = RGBColor.from_string(SLATE)


def _abstract(document: Document, report: dict[str, Any]) -> None:
    """Add abstract, keywords, and the interpretive boundary."""
    _page_break(document)
    _heading(document, "摘要", 1)
    _paragraph(
        document,
        (
            "本文研究财报事件后的市场反应究竟更接近对财务状态（Level）、状态变化"
            "（Delta），还是相对公告前条件预期的新增信息（Innovation）的定价。研究以"
            "Company × Filing Event × Information Set 为原子单位，使用 SEC EDGAR 的"
            "原始 10-K/10-Q 披露、申报接受时间戳、点时 XBRL 财务事实，以及逐证券的"
            "未调整开收盘价、拆股、现金股息、退市结局与 SPY 基准数据。样本在观察"
            "收益之前预注册，共 58 个候选事件、15 家发行人；49 个事件同时具备完整"
            "Level/Delta/Innovation 与 20 个交易日市场超额收益。Innovation 由严格"
            "扩展窗口的样本外 pooled OLS 毛利率预测误差构造，并以仅来自更早样本外"
            "预测误差的标准差进行标准化。主要推断采用发行人整簇 bootstrap 的 "
            "Spearman IC 与增量 rank 回归。标准化 Innovation 的主样本 IC 为 "
            f"{report['innovation_ic']:+.3f}，95% 置信区间为 "
            f"[{report['innovation_ic_ci_low']:+.3f}, {report['innovation_ic_ci_high']:+.3f}]；"
            "相对 Delta 的 IC 差为 +0.064，区间同样跨越零。分位数组合不呈单调关系，"
            "且排除 ASNS 后 Innovation 方向转负，因此研究的预注册结论为“有效性尚未"
            "建立”。然而，幸存者偏差诊断显示，仅保留回看时仍存活的公司会把平均"
            "市场超额收益抬高 1.81 个百分点。本文的阶段性贡献不是宣称发现新因子，"
            "而是建立了一条能够识别并量化前视、幸存者、公告时点和证券身份偏差的"
            "可审计研究管线。"
        ),
    )
    keywords = document.add_paragraph()
    keywords.paragraph_format.space_before = Pt(8)
    label = keywords.add_run("关键词：")
    _set_run_font(label)
    label.font.bold = True
    values = keywords.add_run(
        "财报事件；点时数据；Innovation；毛利率；前视偏差；幸存者偏差；事件研究"
    )
    _set_run_font(values)

    _callout(
        document,
        "论文结论边界",
        (
            "研究框架通过了数据与时间门，但单一毛利率 Innovation 尚未通过统计与"
            "可交易性门。换言之，当前结果支持继续扩展独立数据，不支持直接训练买卖模型。"
        ),
        fill=PALE_RED,
        accent=RED,
    )


def _introduction(document: Document, figures: dict[str, Path]) -> None:
    """Add the motivation, research question, and contribution."""
    _heading(document, "1. 引言", 1)
    _paragraph(
        document,
        (
            "金融市场会在连续时间中吸收历史信息，但投资者真正下达订单的时刻是离散的。"
            "因此，事件驱动决策面对的并不是“公司毛利率是多少”这一静态问题，而是“本次"
            "披露相对于公告前的条件预期新增了什么”。这一视角把财报研究拆成两个输入层："
            "State Features 描述公告前已经形成的状态，Innovation Features 描述新观察值"
            "与即时事前预期之间的偏离。"
        ),
    )
    _picture(document, figures["concept"])
    _source(
        document,
        "该图是本项目的研究假设，而不是经验结论；Innovation 的定义依赖可审计的事前预测。",
    )
    _paragraph(
        document,
        (
            "传统财报因子研究已经证明会计信息与证券价格存在系统关系，并记录了公告后"
            "漂移等现象（Ball and Brown, 1968; Bernard and Thomas, 1989）。但“surprise”"
            "如何定义会显著改变结果：分析师预测与时间序列预测可以捕捉不同的信息集"
            "（Livnat and Mendenhall, 2006）。这使得财报因子验证的主要困难不只是模型"
            "选择，还包括历史上当时真正可得的数据、证券是否后来消失、公告后何时可以"
            "成交，以及研究者是否在看见收益后改变样本或窗口。"
        ),
    )
    _paragraph(
        document,
        (
            "本文的主问题是：在完全相同的点时完整样本上，标准化样本外毛利率 "
            "Innovation 对公告后 20 个交易日的市场超额收益，是否比简单的毛利率同比"
            "变化 Delta 具有更强且更稳定的解释力？我们把这个问题放在预注册、不可变"
            "数据版本和发行人聚类推断之下，以避免事后选择。"
        ),
    )

    _heading(document, "1.1 本文的阶段性贡献", 2)
    _numbered(
        document,
        (
            "概念上，将 Level、Delta 与 model-based Innovation 明确拆分，"
            "并证明 Innovation = Delta + 预期调整，不把 Innovation 错称为分析师 surprise。"
        ),
    )
    _numbered(
        document,
        (
            "数据上，建立从 SEC accession 与 acceptance timestamp 到原始 XBRL 事实、"
            "历史证券身份、公司行动、退市结局和事件收益的点时链条。"
        ),
    )
    _numbered(
        document,
        "方法上，在观察收益前冻结样本、窗口、统计量和判定词汇；对同一发行人的多个事件使用整簇重采样。",
    )
    _numbered(
        document,
        (
            "结果上，拒绝把方向略正但区间跨零的估计包装成有效因子，"
            "同时量化出 1.81 个百分点的幸存者样本抬升。"
        ),
    )


def _literature_and_hypotheses(document: Document) -> None:
    """Position the study and state the frozen hypotheses."""
    _heading(document, "2. 研究背景与可证伪假设", 1)
    _heading(document, "2.1 从信息含量到公告后反应", 2)
    _paragraph(
        document,
        (
            "Ball and Brown（1968）以经验方式连接会计收益与证券价格反应；Bernard and "
            "Thomas（1989）进一步系统研究公告后漂移。本文不试图在 49 个事件上复制整套"
            "PEAD 文献，也不把毛利率变化等同于 SUE。我们的目标更窄：检验一种严格点时"
            "构造的毛利率 Innovation 是否比 Level 和 Delta 更接近事件后的横截面排序。"
        ),
    )
    _paragraph(
        document,
        (
            "Livnat and Mendenhall（2006）表明，基于分析师预测与基于时间序列预测的盈余"
            "surprise 会产生不同漂移强度。这一结果直接提示本研究的局限：本文第一版"
            "Innovation 使用会计历史和滞后行业信息，而未纳入分析师一致预期、公司指引"
            "或文本语义。因此，它是 model-based accounting innovation，不是 Street "
            "surprise。"
        ),
    )
    _paragraph(
        document,
        (
            "偏差控制同样是研究对象。Brown et al.（1992）展示了幸存者截断如何制造"
            "表面可预测性；White（2000）讨论了反复使用同一数据进行模型选择所产生的"
            "data-snooping 风险。本文据此把后来退市证券保留在点时样本中，并在读取收益"
            "之前冻结主要窗口和判定规则。"
        ),
    )

    _page_break(document)
    _heading(document, "2.2 预注册假设", 2)
    _table(
        document,
        "表 1｜研究假设与判定角色",
        ["编号", "假设", "统计角色"],
        [
            ["H1", "标准化 Innovation 与 20 日市场超额收益的 Spearman IC 为正。", "主要确认性"],
            ["H2", "IC(Innovation) - IC(Delta) 为正，且区间排除 0。", "主要对比"],
            ["H3", "控制 Delta 后，Innovation 的 rank 回归系数仍为正。", "必要支持"],
            ["H4", "Innovation 分位数收益随信号强度单调上升。", "经济解释"],
            ["B1", "仅保留存活公司会系统性抬高样本收益。", "偏差诊断"],
        ],
        [900, 6120, 2340],
        font_size=8.5,
    )
    _source(document, "H1-H3、窗口与判定词汇在市场收益接入前冻结；H4 和 B1 用于解释与偏差诊断。")


def _data_section(
    document: Document,
    figures: dict[str, Path],
    prereg: dict[str, Any],
    ledger: dict[str, Any],
    analysis: pd.DataFrame,
) -> None:
    """Describe sources, sample construction, and point-in-time lineage."""
    _heading(document, "3. 数据与样本", 1)
    _heading(document, "3.1 数据来源", 2)
    _table(
        document,
        "表 2｜核心数据源与用途",
        ["数据层", "主要字段", "研究用途", "关键控制"],
        [
            [
                "SEC EDGAR",
                "10-K/10-Q、accession、acceptance timestamp、XBRL facts",
                "构造原始披露与财务期",
                "原始响应 bytes 与 SHA-256；不把 filed date 当可得时点",
            ],
            [
                "历史财务面板",
                "收入、毛利、毛利率、同比滞后、行业滞后",
                "Level、Delta、预测特征",
                "只使用目标公告前已可得事实",
            ],
            [
                "Massive 市场数据",
                "未调整日线开收盘、拆股、现金股息、证券参考",
                "公司总收益与身份映射",
                "价格、公司行动与身份分别过门",
            ],
            [
                "交易日历与 SPY",
                "正常交易日、SPY 未调整价格与股息",
                "执行窗口与市场基准",
                "公司与基准使用同一入场/退出日",
            ],
            [
                "退市结果",
                "active、later delisted、结局时间",
                "保留失败公司",
                "不把未知退市收益默认为 0",
            ],
        ],
        [1500, 2700, 2460, 2700],
        font_size=7.9,
    )
    _source(
        document,
        (
            "SEC API 提供 submissions 与抽取后的 XBRL companyfacts；本文还要求 accession "
            "必须匹配 filing acceptance timestamp。SEC API 说明见 SEC（2024）。"
        ),
    )

    _heading(document, "3.2 样本形成", 2)
    _paragraph(
        document,
        (
            "预注册事件集合包含 58 个候选财报事件、15 家美国发行人，公告接受时间从 "
            "2024-10-30 至 2025-08-14。选择依据是点时特征密度、SEC 身份与市场数据可得性，"
            "不使用事件收益；因此这是技术性试点样本，而不是面向美国股票总体的概率样本。"
            "49 个事件（84.5%）同时具备四种表示和主窗口收益，超过预先计算的最小门槛 47。"
        ),
    )
    _picture(document, figures["pipeline"])
    _source(
        document,
        (
            f"预注册版本 {prereg['dataset_version']}；收益账本版本 {ledger['dataset_version']}。"
            "49 个完整事件中包括 38 个 10-Q、11 个 10-K。"
        ),
    )
    _paragraph(
        document,
        (
            "主样本含 6 个后来退市事件，来自 2 家发行人。收益账本对 58 个事件分别构造 "
            "1、5 和 20 个交易日窗口，共 174 个 event-horizon 观察值；所有 58 个候选事件"
            "的市场超额收益均可计算，主样本的减少完全来自表示特征完整性，而不是收益表现。"
        ),
    )
    _source(
        document,
        (
            f"主样本公告期：{analysis['accepted_at'].min()} 至 {analysis['accepted_at'].max()}；"
            f"主窗口退出日期最晚为 {analysis['exit_at'].max()}。"
        ),
    )


def _methodology(document: Document) -> None:
    """Explain representations, forecasting, returns, and inference."""
    _heading(document, "4. 变量构造与模型", 1)
    _heading(document, "4.1 Level 与 Delta", 2)
    _paragraph(
        document,
        "对公司 i 的财务期 t，毛利率 Level 使用该次原始披露中的毛利与营业收入计算：",
    )
    _equation(document, "Levelᵢ,ₜ = GMᵢ,ₜ = Gross Profitᵢ,ₜ / Revenueᵢ,ₜ")
    _paragraph(
        document,
        (
            "Delta 是相对去年同期、且去年同期事实在当前公告前已经可得的毛利率变化。"
            "滞后值通过 accession 与可得时间连接，而不是简单按行 shift："
        ),
    )
    _equation(document, "Deltaᵢ,ₜ = GMᵢ,ₜ - GMᵢ,ₜ₋₄")

    _heading(document, "4.2 样本外条件预期", 2)
    _paragraph(
        document,
        (
            "第一版条件预期使用扩展窗口 pooled OLS。对每个目标事件，训练集只包含"
            "实际披露可得时间严格早于目标事件的历史观察。相同时间戳的事件作为一个批次"
            "预测，不能互相训练。模型至少需要 200 个历史训练观察。"
        ),
    )
    _equation(document, "GM̂ᵢ,ₜ|ₜ₋ = β̂₀,ₜ₋ + β̂ₜ₋′ Xᵢ,ₜ₋")
    _paragraph(
        document,
        (
            "X 包含四个公告前特征：去年同期毛利率、滞后一季的毛利率同比变化、滞后一季"
            "收入增长，以及通过 as-of 规则构造的滞后行业毛利率中位数。行业统计至少需要"
            "3 家历史可得同组公司；当前季度行业信息被禁止进入模型。"
        ),
    )

    _heading(document, "4.3 Raw 与 Standardized Innovation", 2)
    _equation(document, "Innovation rawᵢ,ₜ = GMᵢ,ₜ - GM̂ᵢ,ₜ|ₜ₋")
    _equation(document, "Innovation stdᵢ,ₜ = Innovation rawᵢ,ₜ / σ̂ₜ₋")
    _paragraph(
        document,
        (
            "σ̂ 只使用目标事件之前已经产生的样本外预测误差，至少需要 100 个历史误差。"
            "同一时点误差、样本内残差和未来误差均被排除。由定义可得："
        ),
    )
    _equation(document, "Innovation raw = Delta + (去年同期毛利率 - 条件预测毛利率)")
    _paragraph(
        document,
        (
            "该恒等式解释了为什么当前 Innovation 仍可能主要是 Delta 的重新表达：只有"
            "括号中的“预期调整”足够有信息，Innovation 才会实质性重排 Delta。"
        ),
    )

    _table(
        document,
        "表 3｜表示变量及其信息含义",
        ["变量", "公式/定义", "含义", "主要失真风险"],
        [
            ["Level", "GM", "当前经营状态", "行业和商业模式截面差异"],
            ["Delta", "GM - GM(t-4)", "相对去年同期改善", "基数、季节性、一次性项目"],
            ["Innovation raw", "GM - OOS Forecast", "相对条件预期的新信息", "预测模型遗漏变量"],
            [
                "Innovation std",
                "raw / prior OOS σ",
                "跨事件可比的标准化新信息",
                "历史误差样本不足或状态变化",
            ],
        ],
        [1450, 2050, 2700, 3160],
        font_size=8.2,
    )
    _source(document, "所有变量均使用原始披露值；latest/restated facts 不进入无偏主信号。")


def _return_and_inference(document: Document) -> None:
    """Describe tradable return construction and statistical inference."""
    _heading(document, "5. 事件收益与统计推断", 1)
    _heading(document, "5.1 可交易收益定义", 2)
    _paragraph(
        document,
        (
            "公告时间不是收益窗口的模糊标签。对每项 SEC filing acceptance，策略在公告之后"
            "第二个未来正常交易日的未调整开盘价进入，并在第 20 个持有交易日未调整收盘价"
            "退出。拆股按生效日调整股数，现金股息保留但不再投资。"
        ),
    )
    _equation(document, "Rᵢ,20 = (Exit wealthᵢ - Entry wealthᵢ) / Entry wealthᵢ")
    _equation(document, "Rᵉˣᶜᵉˢˢᵢ,20 = Rᵢ,20 - RSPY,20")
    _paragraph(
        document,
        (
            "该执行规则故意比“公告日收盘成交”更保守，并确保公司与 SPY 使用同一入场和"
            "退出交易日。本文计算横截面解释力，不把相互重叠的事件直接拼成策略净值，"
            "也不报告年化收益、夏普或交易成本后 P&L。"
        ),
    )

    _heading(document, "5.2 主要统计量", 2)
    _paragraph(
        document,
        "每种表示的主要描述统计为信号与 20 日市场超额收益之间的 Spearman 信息系数：",
    )
    _equation(document, "IC(F) = ρS(rank(Fᵢ), rank(Rᵉˣᶜᵉˢˢᵢ,20))")
    _paragraph(
        document,
        ("主要对比是 IC(Innovation std) - IC(Delta)。增量检验使用 rank 回归："),
    )
    _equation(
        document,
        "rank(Rᵉˣᶜᵉˢˢᵢ,20) = α + βΔ rank(Deltaᵢ) + βI rank(Innovation stdᵢ) + εᵢ",
    )
    _paragraph(
        document,
        (
            "由于同一发行人贡献多个季度事件，10,000 次 bootstrap 以发行人为整簇有放回"
            "抽样，而不是把 49 个事件视为独立。置信水平为 95%，随机种子固定为 20260814。"
            "次要 1 日和 5 日窗口使用 Benjamini-Hochberg 调整；逐发行人剔除、排除 ASNS、"
            "仅 10-Q、按季度分块和单一发行人影响占比均为强制稳健性检查。"
        ),
    )
    _callout(
        document,
        "识别边界",
        (
            "本研究估计的是小型技术样本中的横截面相关与稳定性，不是因果效应。"
            "即便 IC 显著，也仍需交易成本、风险中性化和真正的时间外组合测试。"
        ),
        fill=PALE_GOLD,
        accent=GOLD,
    )


def _results_main(
    document: Document,
    figures: dict[str, Path],
    tables: dict[str, pd.DataFrame],
) -> None:
    """Report primary, incremental, and horizon results."""
    _heading(document, "6. 实证结果", 1)
    _heading(document, "6.1 Level、Delta 与 Innovation 的同样本比较", 2)
    _picture(document, figures["ic"])
    _source(document, "49 个完整事件、15 家发行人；四种表示使用相同 complete-case 样本。")

    labels = {
        "gm_level": "Level",
        "gm_delta_yoy": "Delta",
        "gm_innovation_raw": "Innovation raw",
        "gm_innovation_std": "Innovation std（主信号）",
    }
    ic_rows = []
    for row in tables["representation_ic"].itertuples(index=False):
        ic_rows.append(
            [
                labels[row.feature],
                str(row.events),
                f"{row.spearman_ic:+.3f}",
                f"[{row.ci_low:+.3f}, {row.ci_high:+.3f}]",
            ]
        )
    _table(
        document,
        "表 4｜主要表示的 Spearman IC",
        ["表示", "事件数", "IC", "95% CI"],
        ic_rows,
        [3300, 1500, 1500, 3060],
    )
    _source(document, "所有置信区间均跨越 0；点估计不能单独构成有效性证据。")
    _paragraph(
        document,
        (
            "Level 的 IC 为 -0.153，Delta 接近于零，Raw 与 Standardized Innovation 均约"
            "为 +0.045。Innovation 方向略正，但 95% 区间非常宽，允许从明显负向到明显"
            "正向的多种总体结果。因此 H1 未通过。"
        ),
    )

    _heading(document, "6.2 Innovation 是否比 Delta 多提供信息", 2)
    incremental = tables["incremental"]
    inc_rows = []
    stat_names = {
        "innovation_minus_delta_spearman_ic": "IC(Innovation) - IC(Delta)",
        "incremental_innovation_rank_beta_given_delta": "控制 Delta 后的 Innovation rank beta",
    }
    for row in incremental.itertuples(index=False):
        inc_rows.append(
            [
                stat_names[row.statistic],
                f"{row.point_estimate:+.3f}",
                f"[{row.ci_low:+.3f}, {row.ci_high:+.3f}]",
                f"{row.bootstrap_p_value:.3f}",
            ]
        )
    _table(
        document,
        "表 5｜增量信息检验",
        ["检验", "点估计", "95% CI", "Bootstrap p"],
        inc_rows,
        [4380, 1380, 2280, 1320],
    )
    _source(document, "两项增量检验均无法排除零；H2 与 H3 的确认性证据不足。")

    _picture(document, figures["horizons"])
    _source(document, "5 日窗口差值为 +0.219，但 BH 校正 p=0.340；不能事后替换 20 日主窗口。")
    _paragraph(
        document,
        (
            "5 日 Innovation IC 为 +0.204，显著高于 1 日与 20 日点估计，提示信息反应可能"
            "集中在较短窗口。但该窗口是次要敏感性检验，且多重检验校正后不显著。它只能"
            "形成下一批独立样本的预注册假设，不能改写本轮结论。"
        ),
    )


def _results_robustness(
    document: Document,
    figures: dict[str, Path],
    tables: dict[str, pd.DataFrame],
) -> None:
    """Report economic ordering, quarter stability, and survivorship bias."""
    _heading(document, "6.3 分位数可交易性与季度稳定性", 2)
    _picture(document, figures["quantiles"])
    _source(document, "标准化 Innovation 的 Low/Middle/High 组分别包含 17/16/16 个事件。")
    _paragraph(
        document,
        (
            "低分位的平均 20 日市场超额收益为 -2.89%，中分位为 +0.92%，高分位为 +0.42%。"
            "高分位虽高于低分位，但低于中分位，排序不单调。当前差异更像由负向尾部驱动，"
            "而不是高 Innovation 持续跑赢；因此 H4 未获得支持。"
        ),
    )

    _picture(document, figures["quarters"])
    _source(document, "2024Q4、2025Q1、2025Q2、2025Q3 的事件数分别为 10、10、15、14。")
    _paragraph(
        document,
        (
            "Innovation IC 在 2025Q2 为 -0.420，而其他三个季度为正。这种方向翻转可以"
            "来自宏观状态、行业构成、公司生命周期或纯粹的小样本波动。由于每季只有 "
            "10-15 个事件，本文不对季度差异作结构性解释，只把它记录为外部有效性风险。"
        ),
    )

    _heading(document, "6.4 幸存者偏差", 2)
    _picture(document, figures["survivorship"])
    _source(document, "后来退市证券按事件发生时的真实可投资状态保留；不使用今天的存活股票池回看。")
    survivorship = tables["survivorship"].set_index("sample_id")
    all_mean = float(survivorship.loc["all_events"].mean_market_excess_return)
    survivor_mean = float(survivorship.loc["survivors_only_biased"].mean_market_excess_return)
    inflation = survivor_mean - all_mean
    _paragraph(
        document,
        (
            f"完整 49 事件样本的平均市场超额收益为 {all_mean:.2%}。如果仅保留回看时仍"
            f"存活的 43 个事件，均值变为 {survivor_mean:.2%}，提高 {inflation * 100:.2f} "
            "个百分点；6 个后来退市事件的均值为 -13.54%。这是本轮最明确的偏差结果："
            "即使不改变信号，股票池的回看式清洗也足以改变收益符号。"
        ),
    )
    _callout(
        document,
        "偏差结果与信号结果必须分开",
        (
            "幸存者偏差被清晰量化，并不意味着毛利率 Innovation 已经盈利。前者验证了"
            "研究基础设施的必要性；后者仍需更大、独立的样本。"
        ),
        fill=PALE_TEAL,
        accent=TEAL,
    )


def _repair_and_verdict(
    document: Document,
    figures: dict[str, Path],
    repair: dict[str, Any],
    verdict_checks: pd.DataFrame,
) -> None:
    """Report the ticker-transition repair and locked verdict."""
    _heading(document, "6.5 证券身份修复：OB→TEAD", 2)
    _paragraph(
        document,
        (
            "一项已冻结事件的 20 日持有期跨过 OB→TEAD 证券代码迁移。若只按旧 ticker 抓价，"
            "退出路径会被截断。修复门要求同 CIK、同 provider security id、同 share-class "
            "FIGI 与连续交易日，并检查边界价格尺度、重复 bars 和公司行动。"
        ),
    )
    _picture(document, figures["repair"])
    _source(
        document,
        (
            f"旧代码最后交易日 2025-06-09，新代码首日 {repair['transition_date']}；"
            f"边界隔夜收益 {repair['overnight_return']:.2%}；"
            f"新增 {repair['price_addition_rows']} 条日线。"
        ),
    )
    _paragraph(
        document,
        (
            "补回这项负向事件后，完整样本由 48 增至 49，Innovation IC 从 +0.094 降至 "
            "+0.045。数据质量修复使结果更不利而不是更漂亮，说明研究规则没有迎合正向结论。"
        ),
    )

    _heading(document, "6.6 预注册判定", 2)
    check_names = {
        "data_quality_and_coverage": "数据质量与覆盖门",
        "innovation_ic_positive_and_ci_excludes_zero": "Innovation IC 为正且 CI 排除 0",
        "innovation_minus_delta_ci_excludes_zero": "Innovation-Delta 差异 CI 排除 0",
        "incremental_innovation_beta_positive": "控制 Delta 后 Innovation beta 为正",
        "leave_one_issuer_out_positive_sign_share": "留一发行人正向比例 ≥ 80%",
        "exclude_ASNS_innovation_sign_positive": "排除 ASNS 后 Innovation 仍为正",
        "maximum_single_issuer_absolute_effect_share": "单一发行人影响占比 ≤ 25%",
    }
    check_rows = []
    for row in verdict_checks.itertuples(index=False):
        check_rows.append(
            [
                check_names[row.check_id],
                "通过" if row.passed else "未通过",
                f"{row.observed:+.3f}" if abs(float(row.observed)) < 1 else f"{row.observed:.0f}",
                f"{row.threshold:+.3f}"
                if abs(float(row.threshold)) < 1
                else f"{row.threshold:.0f}",
            ]
        )
    table = _table(
        document,
        "表 6｜冻结判据与观察值",
        ["判据", "状态", "观察值", "阈值"],
        check_rows,
        [5460, 1260, 1320, 1320],
        font_size=8.1,
    )
    for row_index, row in enumerate(verdict_checks.itertuples(index=False), start=1):
        _set_cell_shading(table.cell(row_index, 1), PALE_TEAL if row.passed else PALE_RED)
    _source(document, "7 项判据通过 4 项；终局词汇按预注册顺序判定。")
    _callout(
        document,
        "正式结论：effectiveness_not_established",
        (
            "Innovation 的点估计略正，但主要置信区间跨越零、相对 Delta 的优势不显著、"
            "分位数不单调，并在排除 ASNS 后转负。当前不能把该信号描述为有效，更不能"
            "据此建立买卖建议。"
        ),
        fill=PALE_RED,
        accent=RED,
    )


def _discussion(document: Document) -> None:
    """Interpret the result and separate useful infrastructure from weak signal."""
    _heading(document, "7. 讨论：结果代表什么", 1)
    _heading(document, "7.1 对毛利率 Innovation 的解释", 2)
    _paragraph(
        document,
        (
            "第一，当前结果不是“反证毛利率信息永远无用”，而是说明这一定义、这一预测"
            "信息集和这一小型样本不足以建立稳定解释力。Innovation 的 95% 区间宽，意味着"
            "统计不确定性远大于点估计。"
        ),
    )
    _paragraph(
        document,
        (
            "第二，Innovation 可能仍主要是 Delta 的重新表达。第一版 pooled OLS 使用历史"
            "财务与滞后行业特征，但缺少分析师一致预期、公司指引、订单/价格信息和公告文本。"
            "如果公告前市场预期已包含这些信息，则会计历史模型的预测误差不等于市场真正"
            "接收到的 surprise。"
        ),
    )
    _paragraph(
        document,
        (
            "第三，5 日窗口的较强点估计提示反应可能短于 20 日，但这只是探索性线索。"
            "真正严谨的做法不是在当前样本上把持有期改成 5 日，而是在新样本上提前锁定"
            "5 日假设并复核。"
        ),
    )

    _heading(document, "7.2 对研究框架的解释", 2)
    _paragraph(
        document,
        (
            "就交易信号而言，本轮结果不够好；就研究系统而言，本轮结果是有价值的。"
            "系统成功阻止了三类常见错误：使用未来修订值、删除后来失败的证券、以及因"
            "ticker 迁移而静默删除负向事件。尤其是幸存者样本抬升 1.81 个百分点和 "
            "OB→TEAD 修复后 IC 下降，说明数据偏差的经济量级足以覆盖小型信号本身。"
        ),
    )
    _paragraph(
        document,
        (
            "因此，本项目最独特的资产不是当前某一个因子，而是把“事件当时知道什么、"
            "当时能交易什么、后来谁从数据库消失”转化为可计算、可审计、可复现的研究"
            "对象。这个基础设施可以复用于 EPS surprise、收入 surprise、指引变化、文本"
            "创新和未来的多资产事件模型。"
        ),
    )


def _limitations(document: Document) -> None:
    """State the study's internal and external validity limitations."""
    _heading(document, "8. 局限性", 1)
    limitations = [
        (
            "样本规模有限。49 个完整事件来自 15 家发行人，且同一发行人多次出现；"
            "即使使用整簇 bootstrap，估计仍可能不稳定。"
        ),
        (
            "样本不具总体代表性。发行人因特征密度和市场数据 readiness 进入技术试点，"
            "而非按美国上市公司总体概率抽样。"
        ),
        "时间范围较短。公告集中在 2024Q4-2025Q3，无法覆盖完整经济周期和不同利率、通胀与波动环境。",
        (
            "预期模型仍初级。pooled OLS 没有纳入分析师预测、公司指引、文本语义或"
            "高频市场预期，因此 Innovation 可能测量误差较大。"
        ),
        "基准仍较粗。SPY 能控制同期市场方向，但没有完成点时行业基准、风格因子、beta 或风险中性化。",
        "尚未形成策略回测。本文没有组合形成、头寸重叠、交易成本、做空可得性、容量、滑点与税务处理。",
        (
            "修订偏差在本研究中主要通过原始值契约被阻断，但大样本的 latest/restated "
            "对照尚未作为本次收益检验的独立估计量展开。"
        ),
    ]
    for number, limitation in enumerate(limitations, start=1):
        _manual_numbered(document, number, limitation)
    _callout(
        document,
        "外部沟通建议",
        (
            "向专业人士介绍本稿时，应使用“阶段性方法与数据验证”“有效性尚未建立”"
            "和“偏差量级已被量化”等措辞，避免使用“发现因子”“可交易 alpha”或“策略"
            "收益”等尚未被证据支持的表述。"
        ),
        fill=PALE_GOLD,
        accent=GOLD,
    )


def _next_steps(document: Document) -> None:
    """Specify the independent expansion and eventual route to live research."""
    _heading(document, "9. 下一步研究计划", 1)
    _heading(document, "9.1 先增加独立证据", 2)
    _paragraph(
        document,
        (
            "下一轮的边际价值主要来自扩样，而不是继续在 49 个事件上调模型。当前样本应"
            "完整冻结为开发/发现样本；新增公司和事件不得依据本轮收益选择。一个可操作"
            "目标是把完整事件扩大到当前的 3-5 倍，并显著增加独立发行人、行业、中小盘"
            "与退市/清退案例。该数量是工程目标，不是事后伪装的统计功效承诺。"
        ),
    )
    _heading(document, "9.2 将 5 日窗口变成新的预注册复核", 2)
    _paragraph(
        document,
        (
            "在完全独立的数据批次中，将 5 日市场超额收益预先锁定为复核窗口，同时保留"
            "20 日作为本轮不变的历史主窗口。只有当 Innovation 的方向、置信区间、相对"
            "Delta 的增量和分位数单调性在新样本中共同改善，才升级证据等级。"
        ),
    )
    _heading(document, "9.3 升级预期信息集", 2)
    _table(
        document,
        "表 7｜后续数据与模型扩展",
        ["层级", "新增数据", "目的", "进入条件"],
        [
            [
                "预期",
                "分析师一致预期、预测分散度、公司指引",
                "逼近市场公告前真实条件预期",
                "严格 as-of 快照",
            ],
            [
                "文本",
                "10-Q/10-K、8-K、业绩会文本、语义变化",
                "捕捉数字之外的 innovation",
                "公告时间与版本可审计",
            ],
            [
                "市场",
                "点差、成交量、隔夜跳空、短卖可得性",
                "估计交易摩擦与容量",
                "历史证券身份完整",
            ],
            [
                "风险",
                "行业、规模、价值、动量、利率、波动 regime",
                "区分状态依赖与风险暴露",
                "点时因子与行业映射",
            ],
            [
                "期货",
                "连续合约、换月、期限结构、持仓与宏观事件",
                "扩展到无杠杆期货配置",
                "独立的合约级 PIT 账本",
            ],
        ],
        [1200, 2850, 3000, 2310],
        font_size=8.1,
    )
    _source(document, "新增数据只能向前扩展信息集，不能回填当前已冻结样本的“当时预期”。")

    _heading(document, "9.4 模型复杂度与停止规则", 2)
    _paragraph(
        document,
        (
            "在样本与标签成熟之前，优先使用线性模型、正则化回归和树模型，因为它们更"
            "容易诊断时间泄漏、行业暴露与尾部依赖。深度学习只有在事件数量、文本/序列"
            "输入和时间外验证规模足够，并且稳定胜过透明基线时才有必要。"
        ),
    )
    _paragraph(
        document,
        (
            "建议设置停止规则：如果独立扩样达到至少当前规模的 3 倍后，Innovation 主 IC "
            "区间仍跨零、相对 Delta 无增量且分位数仍不单调，则将“毛利率 Innovation”"
            "降级为无效候选，不再继续调参；但偏差感知的数据与收益账本继续服务于其他"
            "财报和事件信号。"
        ),
    )


def _conclusion(document: Document) -> None:
    """Add the paper conclusion."""
    _heading(document, "10. 结论", 1)
    _paragraph(
        document,
        (
            "本文以 58 个预注册财报事件为起点，在 49 个完整点时事件上比较 Level、Delta "
            "和标准化样本外 Innovation 对 20 个交易日市场超额收益的解释力。Innovation "
            "方向略正，但置信区间跨越零，相对 Delta 的优势不显著，分位数不单调，且对"
            "尾部发行人和季度状态敏感。因此，单一毛利率 Innovation 的有效性尚未建立。"
        ),
    )
    _paragraph(
        document,
        (
            "研究仍然产生了明确价值：仅保留存活证券会把样本平均超额收益抬高 1.81 个"
            "百分点；补回代码迁移导致的负向事件会把 Innovation IC 从 +0.094 降至 +0.045。"
            "这些结果说明，历史数据库偏差的量级可以大于待验证信号本身。"
        ),
    )
    _paragraph(
        document,
        (
            "因此，项目应继续，但研究重点应从“调好一个因子”转向“扩大独立样本、升级"
            "公告前预期信息集，并保持点时、退市与执行契约”。当前框架最重要的功能不是"
            "承诺 alpha，而是可靠地区分真实信息创新、数据回填和回测幻觉。"
        ),
    )


def _references(document: Document) -> None:
    """Add the verified primary literature and official data documentation."""
    _page_break(document)
    _heading(document, "参考文献", 1)
    references = [
        (
            "Ball, R., and P. Brown. 1968. “An Empirical Evaluation of Accounting "
            "Income Numbers.” Journal of Accounting Research 6(2): 159-178. "
            "https://doi.org/10.2307/2490232."
        ),
        (
            "Bernard, V. L., and J. K. Thomas. 1989. “Post-Earnings-Announcement "
            "Drift: Delayed Price Response or Risk Premium?” Journal of Accounting "
            "Research 27: 1-36. https://doi.org/10.2307/2491062."
        ),
        (
            "Brown, S. J., W. Goetzmann, R. G. Ibbotson, and S. A. Ross. 1992. "
            "“Survivorship Bias in Performance Studies.” Review of Financial Studies "
            "5(4): 553-580. https://doi.org/10.1093/rfs/5.4.553."
        ),
        (
            "Livnat, J., and R. R. Mendenhall. 2006. “Comparing the Post-Earnings "
            "Announcement Drift for Surprises Calculated from Analyst and Time Series "
            "Forecasts.” Journal of Accounting Research 44(1): 177-205. "
            "https://doi.org/10.1111/j.1475-679X.2006.00196.x."
        ),
        (
            "White, H. 2000. “A Reality Check for Data Snooping.” Econometrica 68(5): "
            "1097-1126. https://doi.org/10.1111/1468-0262.00152."
        ),
        (
            "U.S. Securities and Exchange Commission. 2024. “EDGAR Application "
            "Programming Interfaces (APIs).” https://www.sec.gov/search-filings/"
            "edgar-application-programming-interfaces. Accessed 2026-08-14."
        ),
    ]
    for item in references:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.30)
        paragraph.paragraph_format.first_line_indent = Inches(-0.30)
        paragraph.paragraph_format.space_after = Pt(9)
        paragraph.paragraph_format.line_spacing = 1.20
        run = paragraph.add_run(item)
        _set_run_font(run)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(INK)


def _appendix(
    document: Document,
    reports: dict[str, Any],
    analysis: pd.DataFrame,
) -> None:
    """Add issuer composition and immutable lineage appendices."""
    _heading(document, "附录 A｜主样本发行人构成", 1)
    grouped = (
        analysis.groupby(["cik", "ticker"], sort=True)
        .agg(
            events=("ticker", "size"),
            first_event=("accepted_at", "min"),
            last_event=("accepted_at", "max"),
            later_delisted=("later_delisted_after_formal_exit", "max"),
        )
        .reset_index()
    )
    rows = []
    for row in grouped.itertuples(index=False):
        rows.append(
            [
                row.ticker,
                row.cik,
                str(row.events),
                str(row.first_event)[:10],
                str(row.last_event)[:10],
                "是" if row.later_delisted else "否",
            ]
        )
    _table(
        document,
        "表 A1｜49 个完整事件的发行人分布",
        ["Ticker", "CIK", "事件", "首个公告", "最后公告", "后来退市"],
        rows,
        [1080, 1800, 900, 1800, 1800, 1980],
        font_size=7.9,
    )
    _source(document, "OB 事件的持有窗口跨越至 TEAD，但发行人身份保持不变。")

    _heading(document, "附录 B｜冻结版本与可复现性", 1)
    diagnostic = reports["diagnostic"]
    _table(
        document,
        "表 B1｜核心不可变数据版本",
        ["对象", "版本 / SHA-256"],
        [
            ["预注册", reports["prereg"]["dataset_version"]],
            ["事件收益账本", reports["ledger"]["dataset_version"]],
            ["最终统计诊断", diagnostic["dataset_version"]],
            ["OB→TEAD 修复", reports["repair"]["dataset_version"]],
            ["预注册 manifest", diagnostic["preregistration_manifest_sha256"]],
            ["收益账本 manifest", diagnostic["ledger_manifest_sha256"]],
        ],
        [2460, 6900],
        font_size=7.5,
    )
    _source(
        document,
        (
            "所有核心输出保留 manifest 与 SHA-256；模拟参数或校准参数不会被描述为经验事实。"
            "本文未调用 OpenAI API 生成任何统计结论。"
        ),
    )


def _load_reports_and_tables() -> tuple[dict[str, Any], dict[str, pd.DataFrame]]:
    """Load the frozen reporting inputs and analysis sample."""
    reports, tables = brief._load_data()
    tables["analysis_sample"] = pd.read_parquet(DIAGNOSTIC_DIR / "analysis_sample.parquet")
    return reports, tables


def build_paper() -> Path:
    """Build the academic-style working paper from frozen outputs."""
    reports, tables = _load_reports_and_tables()
    diagnostic = reports["diagnostic"]
    if diagnostic["verdict"] != "effectiveness_not_established":
        raise ValueError("Unexpected locked verdict; refusing to relabel the paper.")
    if int(diagnostic["primary_analysis_events"]) != 49:
        raise ValueError("Unexpected primary sample size.")
    if not bool(reports["repair"]["transition_gate"]):
        raise ValueError("Ticker-transition repair gate must pass before reporting.")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    figures = _build_figures(tables)

    document = Document()
    _configure_document(document)
    _cover(document)
    _abstract(document, diagnostic)
    _introduction(document, figures)
    _literature_and_hypotheses(document)
    _data_section(
        document,
        figures,
        reports["prereg"],
        reports["ledger"],
        tables["analysis_sample"],
    )
    _methodology(document)
    _return_and_inference(document)
    _results_main(document, figures, tables)
    _results_robustness(document, figures, tables)
    _repair_and_verdict(document, figures, reports["repair"], tables["verdict_checks"])
    _discussion(document)
    _limitations(document)
    _next_steps(document)
    _conclusion(document)
    _references(document)
    _appendix(document, reports, tables["analysis_sample"])

    props = document.core_properties
    props.title = "财报事件信号的点时验证：Level、Delta 与 Innovation 的偏差感知研究"
    props.subject = "Bias-aware point-in-time filing-event effectiveness study"
    props.author = "Filing-to-Signal Lab"
    props.keywords = "SEC EDGAR, point-in-time, Innovation, survivorship bias, event study"
    props.comments = "Working paper draft for professional review; not investment advice."
    document.save(DOCX_PATH)
    return DOCX_PATH


def main() -> None:
    """CLI entry point."""
    build_paper()
    print("ACADEMIC PAPER DOCX DONE")


if __name__ == "__main__":
    main()
